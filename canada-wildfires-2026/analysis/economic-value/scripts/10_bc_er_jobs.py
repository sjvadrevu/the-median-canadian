"""Fallback jobs analysis while Statistics Canada's site is down: BC Stats' republication of the Labour Force Survey
(BC Data Catalogue, 'Labour Force Survey (LFS) Monthly and Annual', lfs_data_tables.xlsx sheet BCER = employment by BC
economic region, annual averages 2016-2025 and 3-month moving averages to September 2026; and
2025_lfs_employment_by_industry.xlsx = BC employment by detailed industry, annual 1987-2025).

Design (difference-in-differences style, descriptive, seven regions only):
  treated  = the five interior and northern BC ERs (Thompson-Okanagan, Kootenay, Cariboo, North Coast and Nechako,
             Northeast): 2.0% to 15.3% of land burned 2023-2026 and 3% to 11% of 2021 employment in forestry
  control  = Vancouver Island and Coast (0.3% burned) and Lower Mainland-Southwest (5.0% burned, almost all of it in
             thinly populated Squamish-Lillooet and the Fraser Canyon; 0.8% of employment in forestry)
  pre      = 2016-2019 annual average (pre-COVID); post = 2023-2025 annual average; 2026 = Jan-Sep 3MMA mean vs the
             same months of 2025.
  DiD      = (post/pre - 1) treated minus (post/pre - 1) control, employment-weighted groups and per region.
LFS ER estimates have large sampling error in small regions (Northeast about 36,000 employed; BC Stats suppresses some
unemployment rates as 'x'). North Coast and Nechako are published together, so their burned shares are combined.
Outputs: bc_er_employment.csv, bc_er_did.csv, bc_forestry_employment.csv"""
import numpy as np
import pandas as pd
import openpyxl
from vcommon import RAW, OUT

wb = openpyxl.load_workbook(RAW / "bcstats" / "bc_lfs_monthly_tables.xlsx", read_only=True)
rows = list(wb["BCER"].iter_rows(values_only=True))
names = [x for x in rows[1] if x]
cols = [i for i, x in enumerate(rows[1]) if x]
ann, mon = [], []
for r in rows[3:]:
    if isinstance(r[1], int) and r[0] is None:
        for n, i in zip(names, cols):
            ann.append(dict(er=n, year=r[1], emp_k=r[i]))
    elif isinstance(r[0], int) and isinstance(r[1], str):
        for n, i in zip(names, cols):
            mon.append(dict(er=n, year=r[0], month=r[1], emp_k=r[i]))
A = pd.DataFrame(ann); M = pd.DataFrame(mon)
A.to_csv(OUT / "bc_er_employment.csv", index=False)
M.to_csv(OUT / "work" / "bc_er_employment_3mma.csv", index=False)

eb = pd.read_csv(OUT / "er_burned_by_year.csv")
eb = eb[eb.PRUID.astype(str) == "59"]
mapn = {"5910": "VANCOUVER ISLAND AND COAST", "5920": "LOWER MAINLAND-SOUTHWEST", "5930": "THOMPSON-OKANAGAN", "5940": "KOOTENAY",
        "5950": "CARIBOO", "5960": "NORTHCOAST AND NECHAKO", "5970": "NORTHCOAST AND NECHAKO", "5980": "NORTHEAST"}
eb["er"] = eb.ERUID.astype(str).map(mapn)
u = eb[eb.year == "2023_2026"].groupby("er").agg(burn=("area_scaled_ha", "sum"), land=("land_km2", "sum"))
u["share_2023_2026"] = u.burn / (u.land * 100)
pre_b = eb[eb.year.isin([str(y) for y in range(2016, 2023)])].groupby("er").area_scaled_ha.sum() / (u.land * 100)
u["share_burned_2016_2022_sum"] = pre_b
cj = pd.read_csv(OUT / "er_census_jobs_2021.csv", dtype={"ERUID": str})
cj["er"] = cj.ERUID.map(mapn)
f = cj.dropna(subset=["er"]).groupby("er")[["employed_all", "forestry_jobs"]].sum()
u["forestry_share_2021"] = f.forestry_jobs / f.employed_all

P = A.pivot(index="er", columns="year", values="emp_k")
out = []
for er in u.index:
    pre = P.loc[er, 2016:2019].mean(); post = P.loc[er, 2023:2025].mean()
    m25 = M[(M.er == er) & (M.year == 2025) & (M.month.isin(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep"]))].emp_k.mean()
    m26 = M[(M.er == er) & (M.year == 2026)].emp_k.mean()
    out.append(dict(er=er, share_land_burned_2023_2026=u.loc[er, "share_2023_2026"], share_burned_2016_2022_sum=u.loc[er, "share_burned_2016_2022_sum"],
                    forestry_share_2021=u.loc[er, "forestry_share_2021"], emp_pre_2016_2019_k=pre, emp_post_2023_2025_k=post,
                    change_post_vs_pre=post / pre - 1, emp_janSep2025_3mma_k=m25, emp_janSep2026_3mma_k=m26, change_2026_vs_2025=m26 / m25 - 1))
D = pd.DataFrame(out)
D["group"] = np.where(D.er.isin(["VANCOUVER ISLAND AND COAST", "LOWER MAINLAND-SOUTHWEST"]), "control (metro coast)", "treated (interior and north)")
g = D.groupby("group")[["emp_pre_2016_2019_k", "emp_post_2023_2025_k", "emp_janSep2025_3mma_k", "emp_janSep2026_3mma_k"]].sum()
g["change_post_vs_pre"] = g.emp_post_2023_2025_k / g.emp_pre_2016_2019_k - 1
g["change_2026_vs_2025"] = g.emp_janSep2026_3mma_k / g.emp_janSep2025_3mma_k - 1
did = g.loc["treated (interior and north)", "change_post_vs_pre"] - g.loc["control (metro coast)", "change_post_vs_pre"]
did26 = g.loc["treated (interior and north)", "change_2026_vs_2025"] - g.loc["control (metro coast)", "change_2026_vs_2025"]
D["did_vs_control_post"] = D.change_post_vs_pre - g.loc["control (metro coast)", "change_post_vs_pre"]
D["did_vs_control_2026"] = D.change_2026_vs_2025 - g.loc["control (metro coast)", "change_2026_vs_2025"]
D.to_csv(OUT / "bc_er_did.csv", index=False)
g.to_csv(OUT / "work" / "bc_er_did_groups.csv")

# BC detailed industry
ws = openpyxl.load_workbook(RAW / "bcstats" / "bc_lfs_industry_2025.xlsx", read_only=True)["naicsann"]
rr = list(ws.iter_rows(values_only=True))
ix = [i for i, x in enumerate(rr[2]) if x is not None]
yrs = [rr[2][i] for i in ix]
keep = {}
for r in rr:
    if r[0] and any(k in str(r[0]) for k in ["All Industries", "Forestry and logging", "Wood Product", "Paper Manu", "Accommodation Services"]):
        keep[str(r[0]).strip()] = [r[i] for i in ix]
F = pd.DataFrame(keep, index=yrs)
F["forestry_wood_paper_k"] = F.iloc[:, 1] + F.iloc[:, 2] + F.iloc[:, 3]
F.to_csv(OUT / "bc_forestry_employment.csv")
pd.set_option("display.width", 220)
print(D.round(3).to_string()); print(g.round(4)); print("DiD post vs pre:", round(did, 4), "DiD 2026 vs 2025:", round(did26, 4))
print(F.loc[2015:].round(1).to_string())
