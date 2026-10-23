"""Second round of people costs (2023 to 2026): health care rebuilt for 2024-2026, physician fees, Quebec unit costs,
evacuation gaps, days away, lost-work revisions, EI claims, health system costs of evacuations.

Every quoted number is checked against the saved source file (function `check`); values read from data files are
parsed directly. Dollars go to 2025 CAD with CPI all-items (Bank of Canada Valet V41690973), the bill's basis:
spending by fiscal year counts as the season year (season-year CPI); CIHI and CIHI-type unit prices use fiscal-year CPI.
No death counts or death values are used.

Run after health2_01_aqbat_rebuild.py. Outputs (costs-full/): health2_costs.csv (same columns as health_costs.csv),
health2_health_unit_costs.csv, health2_health_costs_by_province.csv, health2_days_away_events.csv,
health2_lost_wages_by_year.csv, health2_ei_excess_claims.csv
"""
import html
import json
import os
import re
import warnings
import zipfile

import numpy as np
import openpyxl
import pandas as pd

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
CF = os.path.dirname(HERE)
AN = os.path.dirname(CF)
R1 = os.path.join(CF, "raw", "health")
R2 = os.path.join(CF, "raw", "health2")
RG = os.path.join(CF, "raw", "gov")
ACC = "2026-10-22"


def text(path):
    t = open(path, encoding="utf-8", errors="ignore").read()
    if path.endswith((".html", ".htm", ".xml")):
        t = re.sub(r"(?s)<script.*?</script>|<style.*?</style>", " ", t)
        t = html.unescape(re.sub(r"<[^>]+>", " ", t))
    t = t.replace("ﬁ", "fi").replace("‑", "-").replace(" ", " ").replace("\xa0", " ").replace("’", "'")
    return re.sub(r"\s+", " ", t)


_cache = {}


def check(path, *quotes):
    if path not in _cache:
        _cache[path] = text(path)
    for q in quotes:
        q2 = re.sub(r"\s+", " ", q.replace("’", "'"))
        assert q2 in _cache[path], f"not found in {path}: {q}"


# ------------------------------------------------------------------ CPI (same series and conventions as the first round)
obs = json.load(open(os.path.join(AN, "losses", "raw", "cpi", "boc_V41690973.json")))["observations"]
cpi = pd.DataFrame([(o["d"], float(o["V41690973"]["v"])) for o in obs], columns=["d", "v"])
cpi["d"] = pd.to_datetime(cpi.d)
CPI_Y = cpi.groupby(cpi.d.dt.year).v.mean()
C25 = CPI_Y[2025]


def cpi_fy(y):
    return cpi[(cpi.d >= f"{y}-04-01") & (cpi.d <= f"{y + 1}-03-01")].v.mean()


def to25(v, y):
    return v * C25 / CPI_Y[y]


BASIS = "2025 CAD (CPI all-items annual average)"
COLS = ["item", "jurisdiction", "period", "quantity", "unit_cost", "amount_low_cad", "amount_central_cad", "amount_high_cad",
        "dollars_basis", "method", "source_name", "url", "accessed", "overlaps_with", "additive", "notes"]
rows = []


def add(**k):
    base = {c: "" for c in COLS}
    base.update(dollars_basis=BASIS, accessed=ACC, additive="no")
    base.update(k)
    for c in ("amount_low_cad", "amount_central_cad", "amount_high_cad"):
        if base[c] != "":
            base[c] = round(float(base[c]))
    rows.append(base)


H1 = pd.read_csv(os.path.join(CF, "health_costs.csv"))       # first round, read only
PROVS = ["NL", "PE", "NS", "NB", "QC", "ON", "MB", "SK", "AB", "BC", "YT", "NT", "NU"]
csd = pd.read_csv(os.path.join(R2, "census", "csd_pop2021_named.csv"), dtype={"REF_AREA": str}).set_index("REF_AREA")
POP = lambda code: int(csd.loc[code, "OBS_VALUE"])            # Census 2021 population of a census subdivision
HH = pd.read_csv(os.path.join(R2, "census", "csd_hh2021.csv"), usecols=["REF_AREA", "OBS_VALUE"], dtype={"REF_AREA": str}).set_index("REF_AREA").OBS_VALUE

# =================================================================== GAP 4: HEALTH CARE
# ---- counts and provincial shares from health2_01
cnt = pd.read_csv(os.path.join(CF, "health2_aqbat_counts.csv"))
shr = pd.read_csv(os.path.join(CF, "health2_aqbat_province_shares.csv"))
OUT_ED = ["Respiratory emergency room visits", "Cardiac emergency room visits"]
OUT_ADM = {"resp": "Respiratory hospital admissions", "card": "Cardiac hospital admissions"}


def counts(k, y):
    r = cnt[(cnt.outcome == k) & (cnt.year == y)].iloc[0]
    return float(r.low), float(r.central), float(r.high)


def shares(k, y):
    s = shr[(shr.outcome == k) & (shr.year == y)].set_index("prov").share
    return s.reindex(PROVS).fillna(0.0)


# ---- unit cost 1: CIHI direct cost per ED visit by province (Spending Trends data tables 2015-16 to 2024-25, Table 3)
CIHI_NEW = os.path.join(R2, "cihi", "spending-trends-health-service-delivery-2015-2024-data-tables-en.xlsx")
wb = openpyxl.load_workbook(CIHI_NEW, read_only=True, data_only=True)
t3 = list(wb["Table 3"].iter_rows(values_only=True))
ie = [i for i, h in enumerate(t3[1]) if h and "Emergency" in str(h)][0]
NAME2P = {"Newfoundland and Labrador": "NL", "Prince Edward Island": "PE", "Nova Scotia": "NS", "New Brunswick": "NB",
          "Quebec": "QC", "Ontario": "ON", "Manitoba": "MB", "Saskatchewan": "SK", "Alberta": "AB", "British Columbia": "BC",
          "Yukon": "YT", "Northwest Territories": "NT", "Nunavut": "NU", "Canada (excluding Quebec and Nunavut)": "CAN"}
EDC = {}
for r in t3[2:]:
    if r[0] == "Direct cost per patient day/visit" and r[1] in NAME2P and r[2] in ("2023–2024", "2024–2025"):
        v = r[ie]
        EDC[(NAME2P[r[1]], int(r[2][:4]))] = float(v) if isinstance(v, (int, float)) else np.nan
assert EDC[("CAN", 2023)] == 241.6 and EDC[("CAN", 2024)] == 251.9 and EDC[("ON", 2024)] == 229.3, EDC

# ---- unit cost 2: Quebec direct cost per ED visit, MSSS AS-471 (centre 6240 L'urgence), FY2023-24, 25 establishments
AS471 = os.path.join(R2, "qc", "as471_2024")
qrows = []
for f in sorted(os.listdir(AS471)):
    if not f.endswith(".txt"):
        continue
    t = open(os.path.join(AS471, f), encoding="utf-8", errors="ignore").read()
    nm = re.search(r"Nom de l'établissement:\s*(.+)", t)
    for m in re.finditer(r"/ 6240-00", t):
        seg = t[m.start(): m.start() + 9000]
        if "L'urgence" not in seg[:1500]:
            continue

        def nums(label):
            mm = re.search(label + r"([^\n]*)", seg)
            return [float(x.replace(" ", "").replace(" ", "").replace(",", ".")) for x in
                    re.findall(r"\d{1,3}(?:[  ]\d{3})*(?:,\d+)?", mm.group(1))]
        gross = nums(r"COÛTS DIRECTS BRUTS\(L\.12 \+ L\.16\)\s+17")[-1]
        visits = nums(r"TOTAL \(L\.25 à L\.27\)\s+28")[-1]
        unit = nums(r"Coût unitaire brut \(L\.17 - L\.23\) / L\.28\s+29")[-1]
        qrows.append(dict(code=f[:-4], establishment=nm.group(1).strip() if nm else "", gross_direct_cost=gross,
                          visits=visits, unit_cost_gross_less_adjustments=unit))
        break
qd = pd.DataFrame(qrows)
QC_VIS = qd.visits.sum()
QC_ED = float((qd.visits * qd.unit_cost_gross_less_adjustments).sum() / QC_VIS)
assert len(qd) == 25 and abs(QC_ED - 287.38) < 0.01, (len(qd), QC_ED)
check(os.path.join(AS471, "1104-5143.txt"), "Coût unitaire brut (L.17 - L.23) / L.28 29 XXXX 237,33 XXXX 235,21")
EDC[("QC", 2023)] = QC_ED
EDC[("QC", 2024)] = QC_ED * cpi_fy(2024) / cpi_fy(2023)      # no FY2024-25 Quebec value used; CPI-carried
ACTUAL = {k for k, v in EDC.items() if np.isfinite(v)}
for p in PROVS:
    for fy in (2023, 2024):
        if (p, fy) not in ACTUAL:
            if fy == 2024 and (p, 2023) in ACTUAL:
                EDC[(p, fy)] = EDC[(p, 2023)] * cpi_fy(2024) / cpi_fy(2023)      # PE 2024-25 not reported
            else:
                EDC[(p, fy)] = EDC[("CAN", fy)]                                  # NU: Canada excl. QC and NU
# full hospital cost ratio (CIHI 2020: full $304 vs direct $158 in 2018-19; first-round ratio 304/157.6)
FULL = 304 / 157.6
check(os.path.join(R2, "cihi", "cihi_ed_2020_extract.txt"), "full hospital cost of an ED visit to $304 in 2018–2019",
      "The direct cost per ED visit has risen from $96 in 2005–2006 to $158 in 2018–2019")

# ---- unit cost 3: physician fee per ED visit, CIHI National Physician Database 2023-24, Table A.3 (cost per service, FFS)
NPDB = os.path.join(R2, "cihi", "npdb_zip", "npdb-utilization-data-tables-2023-en.xlsx")
a3 = list(openpyxl.load_workbook(NPDB, read_only=True, data_only=True)["Table A.3"].iter_rows(values_only=True))
i0 = [i for i, r in enumerate(a3) if r[0] == "Total physicians" and r[1] == "N.L."][0]
hdr = a3[i0][1:13]
ABB = {"N.L.": "NL", "P.E.I.": "PE", "N.S.": "NS", "N.B.": "NB", "Que.": "QC", "Ont.": "ON", "Man.": "MB", "Sask.": "SK",
       "Alta.": "AB", "B.C.": "BC", "Y.T.": "YT", "Total": "CAN"}
blk = a3[i0 + 1: i0 + 32]
names = [r[0] for r in blk]
j_maj = names.index("Total major assessments"); j_oth = names.index("Total other assessments")
r_maj = [r for r in blk[:j_maj] if r[0] == "Hospital outpatient"][0]
r_oth = [r for r in blk[j_maj + 1:j_oth] if r[0] == "Hospital outpatient"][0]
r_spc = [r for r in blk if r[0] == "Out-of-hours/emergency"][0]
PHY = {}
for j, h in enumerate(hdr):
    p = ABB[h]
    PHY[p] = (float(r_oth[j + 1]), float(r_maj[j + 1]), float(r_maj[j + 1]) + float(r_spc[j + 1]))   # low, central, high
assert PHY["CAN"] == (44.4, 83.51, 83.51 + 106.9), PHY["CAN"]
for p in ("NT", "NU"):
    PHY[p] = PHY["CAN"]

# ---- unit cost 4: CIHI Patient Cost Estimator FY2021-22, hospital and physician cost per typical stay, by jurisdiction
wbp = openpyxl.load_workbook(os.path.join(R1, "patient-cost-estimator-results-en.xlsx"), read_only=True)
pr = list(wbp["Table 1"].iter_rows(values_only=True))
ph = [re.sub(r"\s+", " ", str(h)).strip() for h in pr[1][:14]]
pce = pd.DataFrame([r[:14] for r in pr[2:] if r[0] and str(r[0]).startswith("20")], columns=ph)
pce = pce[pce["Fiscal year"] == "2021–2022"].copy()
pce["vol"] = pd.to_numeric(pce["Volume (typical cases)"].astype(str).str.replace("*", "", regex=False), errors="coerce").fillna(0)
HC_ = "Estimated average hospital cost (all age groups)"; PC_ = "Estimated average physician cost (all age groups)"
pce["phys"] = pd.to_numeric(pce[PC_], errors="coerce")
JUR = {v: k for k, v in NAME2P.items()}
JUR["CAN"] = "Canada"
CMG = {"resp": ["139", "138", "136", "147", "141", "097"], "card": ["196", "202", "194", "204", "208"]}
PCE = {}
for g, codes in CMG.items():
    s = pce[pce["Case Mix Group"].isin(codes)]
    cm = s.groupby(["Jurisdiction", "Case Mix Group"]).agg(vol=("vol", "sum"), cost=(HC_, "first"), phys=("phys", "first")).reset_index()
    phys_pool = cm[cm.phys.notna() & (cm.Jurisdiction != "Canada")]
    pw = float((phys_pool.vol * phys_pool.phys).sum() / phys_pool.vol.sum())
    pmin = float(phys_pool.groupby("Case Mix Group").apply(lambda x: (x.vol * x.phys).sum() / x.vol.sum()).min())
    pmax = float(phys_pool.groupby("Case Mix Group").apply(lambda x: (x.vol * x.phys).sum() / x.vol.sum()).max())
    for p in PROVS + ["CAN"]:
        j = JUR.get(p, "Canada")
        c = cm[cm.Jurisdiction == j] if j in set(cm.Jurisdiction) else cm[cm.Jurisdiction == "Canada"]
        hw = float((c.vol * c.cost).sum() / c.vol.sum())
        cp = c[c.phys.notna()]
        phw = float((cp.vol * cp.phys).sum() / cp.vol.sum()) if len(cp) else pw
        PCE[(g, p)] = dict(hosp=(float(c.cost.min()), hw, float(c.cost.max())),
                           phys=(float(cp.phys.min()) if len(cp) else pmin, phw, float(cp.phys.max()) if len(cp) else pmax),
                           phys_reported=bool(len(cp)))
f_ed = {2023: C25 / cpi_fy(2023), 2024: C25 / cpi_fy(2024)}
f_npdb = C25 / cpi_fy(2023)
f_pce = C25 / cpi_fy(2021)

uc = []
for p in PROVS:
    uc.append(dict(prov=p, ed_direct_fy2023_24=round(EDC[(p, 2023)], 2), ed_direct_fy2024_25=round(EDC[(p, 2024)], 2),
                   ed_full_fy2024_25=round(EDC[(p, 2024)] * FULL, 2), phys_per_ed_visit_low=PHY[p][0], phys_per_ed_visit_central=PHY[p][1],
                   phys_per_ed_visit_high=round(PHY[p][2], 2),
                   resp_stay_hosp_min=round(PCE[("resp", p)]["hosp"][0]), resp_stay_hosp_mean=round(PCE[("resp", p)]["hosp"][1]),
                   resp_stay_hosp_max=round(PCE[("resp", p)]["hosp"][2]), card_stay_hosp_mean=round(PCE[("card", p)]["hosp"][1]),
                   resp_stay_phys_mean=round(PCE[("resp", p)]["phys"][1]), card_stay_phys_mean=round(PCE[("card", p)]["phys"][1]),
                   phys_from_own_jurisdiction=PCE[("resp", p)]["phys_reported"]))
pd.DataFrame(uc).to_csv(os.path.join(CF, "health2_health_unit_costs.csv"), index=False)


def cost_year(y):
    """Hospital (ED + admissions) and physician costs by province, 2025 dollars, (low, central, high)."""
    fy = 2023 if y == 2023 else 2024
    s_ed = shares("Respiratory emergency room visits", y)
    s_ad = shares("Respiratory hospital admissions", y)
    ned = [sum(counts(k, y)[i] for k in OUT_ED) for i in range(3)]
    nad = {g: counts(k, y) for g, k in OUT_ADM.items()}
    out = []
    for p in PROVS:
        e = [ned[i] * s_ed[p] for i in range(3)]
        hosp_ed = (e[0] * EDC[(p, fy)] * f_ed[fy], e[1] * EDC[(p, fy)] * FULL * f_ed[fy], e[2] * EDC[(p, fy)] * FULL * f_ed[fy])
        phys_ed = tuple(e[i] * PHY[p][i] * f_npdb for i in range(3))
        ha = [0.0, 0.0, 0.0]; pa = [0.0, 0.0, 0.0]
        for g in ("resp", "card"):
            a = [nad[g][i] * s_ad[p] for i in range(3)]
            for i in range(3):
                ha[i] += a[i] * PCE[(g, p)]["hosp"][i] * f_pce
                pa[i] += a[i] * PCE[(g, p)]["phys"][i] * f_pce
        out.append(dict(year=y, prov=p, ed_visits_central=e[1], admissions_central=sum(nad[g][1] for g in nad) * s_ad[p],
                        hosp_ed_low=hosp_ed[0], hosp_ed_central=hosp_ed[1], hosp_ed_high=hosp_ed[2],
                        hosp_adm_low=ha[0], hosp_adm_central=ha[1], hosp_adm_high=ha[2],
                        phys_low=phys_ed[0] + pa[0], phys_central=phys_ed[1] + pa[1], phys_high=phys_ed[2] + pa[2]))
    return pd.DataFrame(out)


bypr = pd.concat([cost_year(y) for y in (2023, 2024, 2025, 2026)])
bypr.round(2).to_csv(os.path.join(CF, "health2_health_costs_by_province.csv"), index=False)
tot = bypr.groupby("year").sum(numeric_only=True)

MATZ_SRC = "Matz C.J. et al. (Health Canada and ECCC) 2026, GeoHealth 10(2):e2025GH001565, Table 1 (2023 counts) and Table S.1 (AQBAT coefficients)"
UNIT_SRC = ("CIHI Spending Trends in Health Service Delivery, 2015-2016 to 2024-2025, Data Tables, Table 3 (ED direct cost per visit, "
            "by province, FY2023-24 and FY2024-25); MSSS Quebec, Rapport financier annuel AS-471 2023-2024, centre 6240 L'urgence, "
            "25 establishments; CIHI Patient Cost Estimator Results (2024), FY2021-22, by jurisdiction; full-cost ratio 304/157.6 (CIHI 2020)")
UNIT_URL = ("https://www.cihi.ca/sites/default/files/document/spending-trends-health-service-delivery-2015-2024-data-tables-en.xlsx ; "
            "https://publications.msss.gouv.qc.ca/msss/fichiers/statistiques/rapports-financiers/2024/<code>.pdf ; "
            "https://www.cihi.ca/sites/default/files/document/patient-cost-estimator-results-en.xlsx")
PHYS_SRC = ("CIHI National Physician Database 2023-2024, Utilization data tables, Table A.3 (FFS cost per service, total physicians, "
            "major and other assessments in hospital outpatient settings, special calls out-of-hours/emergency); CIHI Patient Cost "
            "Estimator FY2021-22 (estimated average physician cost per typical stay, AB BC MB NL NS ON SK)")
PHYS_URL = ("https://www.cihi.ca/sites/default/files/document/npdb-data-tables-2023-en.zip ; "
            "https://www.cihi.ca/sites/default/files/document/patient-cost-estimator-results-en.xlsx")

# 2023: official counts; re-cost and add physician fees
h1_23 = H1[H1.period.eq("2023 (May-Sep)") & H1.item.str.startswith("Health care: extra respiratory and cardiac")]
old = [h1_23[c].astype(float).sum() for c in ("amount_low_cad", "amount_central_cad", "amount_high_cad")]
new = [tot.loc[2023, f"hosp_ed_{b}"] + tot.loc[2023, f"hosp_adm_{b}"] for b in ("low", "central", "high")]
add(item="Health care: 2023 hospital cost adjustment for provincial unit costs (Quebec AS-471, CIHI provincial ED and stay costs)",
    jurisdiction="Canada", period="2023 (May-Sep)", quantity="Health Canada 2023 counts (2,310 ED visits, 810 admissions) allocated to provinces by our exposure",
    unit_cost=f"ED direct cost per visit by province FY2023-24 (QC {QC_ED:.2f} from AS-471; Canada excl. QC/NU 241.6); stays by province (PCE FY2021-22)",
    amount_low_cad=new[0] - old[0], amount_central_cad=new[1] - old[1], amount_high_cad=new[2] - old[2],
    method="re-costed 2023 hospital cost (province shares of our AQBAT exposure x provincial unit costs, Quebec from MSSS) minus the first-round 2023 rows, which used one national cost; low/central/high conventions as first round",
    source_name=UNIT_SRC + "; " + MATZ_SRC, url=UNIT_URL, overlaps_with="adjusts health_costs.csv 2023 Health Canada-based ED and admission rows (add this increment to them)",
    additive="yes", notes=f"New 2023 hospital total {new[1] / 1e6:.2f} M central (first round {old[1] / 1e6:.2f} M). Quebec ED direct cost {QC_ED:.2f} per visit (3.40 M visits, 25 establishments) vs 241.6 for Canada excl. QC/NU.")
add(item="Health care: physician fees for smoke-related ED visits and hospital stays (Health Canada counts)", jurisdiction="Canada",
    period="2023 (May-Sep)", quantity="2,310 ED visits (1,520 to 3,190); 810 admissions (480 to 1,140)",
    unit_cost="ED: Canada FFS cost per hospital-outpatient major assessment 83.51 (other assessment 44.40; plus out-of-hours special call 106.90), FY2023-24, provincial values used; stay: PCE physician cost per typical stay",
    amount_low_cad=tot.loc[2023, "phys_low"], amount_central_cad=tot.loc[2023, "phys_central"], amount_high_cad=tot.loc[2023, "phys_high"],
    method="ED: one physician assessment per visit (low: other assessment; central: major assessment; high: major assessment plus out-of-hours special call), by province; stays: PCE average physician cost per typical case for the same CMGs (low cheapest, high dearest)",
    source_name=PHYS_SRC + "; " + MATZ_SRC, url=PHYS_URL, overlaps_with="none: CIHI hospital costs exclude physician pay, so first-round rows exclude it",
    additive="yes", notes="ESTIMATE of physician pay. CIHI publishes no physician cost per ED visit; NPDB cost per service covers fee-for-service only (ED doctors on alternative payment plans are assumed to cost the same per service). NT and NU use the Canada value.")

h1_scaled = H1[H1.item.str.startswith("Health care: smoke-related emergency visits and admissions, scaled")]
_bt = pd.read_csv(os.path.join(CF, "health2_aqbat_backtest.csv"))
_bt = _bt[_bt.outcome.isin(OUT_ED + list(OUT_ADM.values()))]
BT_MIN, BT_MAX = float(_bt.hc_over_predicted.min()), float(_bt.hc_over_predicted.max())
for y in (2024, 2025, 2026):
    old_y = h1_scaled[h1_scaled.period.str.startswith(str(y))].iloc[0]
    ned = [sum(counts(k, y)[i] for k in OUT_ED) for i in range(3)]
    nad = [sum(counts(k, y)[i] for k in OUT_ADM.values()) for i in range(3)]
    rr = cnt[(cnt.outcome == "Respiratory emergency room visits") & (cnt.year == y)].iloc[0]
    add(item="Health care: smoke-related ED visits and hospital stays, rebuilt with AQBAT and our exposure (hospital cost)",
        jurisdiction="Canada (exposure: all provinces with monitors; MB, NL, NT imputed in 2025-26)", period=f"{y} (May-Sep)",
        quantity=f"{ned[1]:.0f} ED visits ({ned[0]:.0f} to {ned[2]:.0f}); {nad[1]:.0f} admissions ({nad[0]:.0f} to {nad[2]:.0f}); ratio to 2023 {rr.ratio_to_2023_central:.3f}",
        unit_cost="ED: CIHI FY2024-25 direct cost by province (QC: AS-471 FY2023-24 at FY2024-25 prices), full = x1.929; stays: PCE by province",
        amount_low_cad=tot.loc[y, "hosp_ed_low"] + tot.loc[y, "hosp_adm_low"],
        amount_central_cad=tot.loc[y, "hosp_ed_central"] + tot.loc[y, "hosp_adm_central"],
        amount_high_cad=tot.loc[y, "hosp_ed_high"] + tot.loc[y, "hosp_adm_high"],
        method="counts = Health Canada 2023 count x S(y)/S(2023), S = sum over DAs of age-weighted population x coverage x (1 - exp(-beta x May-Sep mean increment)), AQBAT betas (Table S.1); range = HC 2023 95% interval combined in quadrature with the 2019-2022 back-test error; costs by province (shares of S) x provincial unit costs; low = low count x direct cost / cheapest CMG, central and high = full cost, high count x dearest CMG",
        source_name=MATZ_SRC + "; analysis/health exposure (R50 increments); " + UNIT_SRC, url="costs-full/health2_aqbat_counts.csv ; " + UNIT_URL,
        overlaps_with=f"REPLACES health_costs.csv row 'Health care: smoke-related emergency visits and admissions, scaled from Health Canada 2023' for {y} (central {float(old_y.amount_central_cad) / 1e6:.2f} M; it used a death-based ratio, now excluded): drop that row when adding this one; overlaps our own all-cause ED estimate (do not add both)",
        additive="yes", notes=f"ESTIMATE, not an official count. Back-test of the same method on 2019-2022 against Health Canada's own numbers (ED and admissions): HC/predicted {BT_MIN:.2f} to {BT_MAX:.2f} (over-predicts the light years 2019 and 2022, under-predicts 2021). Baseline rates uniform across provinces. 2026 smoke to 2026-10-19 (May-Sep complete).")
    add(item="Health care: physician fees for smoke-related ED visits and hospital stays (rebuilt counts)", jurisdiction="Canada",
        period=f"{y} (May-Sep)", quantity=f"{ned[1]:.0f} ED visits; {nad[1]:.0f} admissions",
        unit_cost="as the 2023 physician row (NPDB FY2023-24, PCE FY2021-22)",
        amount_low_cad=tot.loc[y, "phys_low"], amount_central_cad=tot.loc[y, "phys_central"], amount_high_cad=tot.loc[y, "phys_high"],
        method="rebuilt counts by province x physician cost per ED assessment and per typical stay (low/central/high as the 2023 row)",
        source_name=PHYS_SRC, url=PHYS_URL, overlaps_with="none (physician pay is outside CIHI hospital costs)", additive="yes",
        notes="ESTIMATE (counts and fee basis).")

# other AQBAT outcomes (reference counts, no health-care dollars) and the welfare valuation
VAL = {"Respiratory emergency room visits": 3480, "Cardiac emergency room visits": 7656, "Child acute bronchitis episodes": 543,
       "Asthma symptom days": 89.4, "Restricted activity days": 83.5, "Acute respiratory symptom days": 22.6}
SI = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", zipfile.ZipFile(os.path.join(AN, "health", "raw", "crf",
             "GH2-10-e2025GH001565-s001.docx")).read("word/document.xml").decode("utf-8"))))
for frag in ["3480", "7656", "543", "89.4", "83.5", "22.6"]:
    assert re.search(r"\$\s*" + re.escape(frag), SI), frag
v23 = sum(VAL[k] * counts(k, 2023)[1] for k in VAL)
for y in (2024, 2025, 2026):
    q = "; ".join(f"{k.lower()} {counts(k, y)[1]:,.0f} ({counts(k, y)[0]:,.0f} to {counts(k, y)[2]:,.0f})" for k in
                  ["Asthma symptom days", "Child acute bronchitis episodes", "Restricted activity days", "Acute respiratory symptom days"])
    add(item="Health Canada outcomes without health-care dollars: rebuilt counts (reference)", jurisdiction="Canada", period=f"{y} (May-Sep)",
        quantity=q, dollars_basis="", method="same anchoring as the ED rows; age weights 5-19 (asthma, bronchitis) and 5+ (restricted activity, symptom days)",
        source_name=MATZ_SRC, url="costs-full/health2_aqbat_counts.csv", additive="no", notes="ESTIMATE; counts only.")
    ratio_v = sum(VAL[k] * counts(k, y)[1] for k in VAL) / v23
    add(item="Health Canada-method acute morbidity valuation, rebuilt (welfare value, not health care spending)", jurisdiction="Canada",
        period=f"{y} (May-Sep)", quantity=f"ratio to 2023 {ratio_v:.3f} (value-weighted over outcomes)",
        amount_low_cad=to25(150e6, 2023) * ratio_v, amount_central_cad=to25(820e6, 2023) * ratio_v, amount_high_cad=to25(2.0e9, 2023) * ratio_v,
        method="Health Canada 2023 valuation ($820M [$150M to $2.0B], 2023 CAD) x value-weighted ratio of rebuilt counts (Table S.2 unit values)",
        source_name=MATZ_SRC + "; Table S.2", url="https://doi.org/10.1029/2025GH001565",
        overlaps_with="health care rows (medical costs inside it); lost work (restricted activity days)", additive="no",
        notes="Not spending; for context only, like the first-round 2023 reference row.")

# =================================================================== GAP 2: EVACUATION GAPS
# Alberta 2025 Emergency Evacuation Payments
ALSS = os.path.join(R2, "ab", "alss-annual-report-2025-2026.txt")
check(ALSS, "Approximately 5,700 Albertans were required to evacuate their homes in 2025 due to wildfires.",
      "provided $4.6 million in Emergency Evacuation Payments to over 4,500 Albertans, including nearly 1,400 children.")
add(item="Evacuations: Alberta Emergency Evacuation Payments", jurisdiction="AB", period="2025 (FY2025-26)",
    quantity="over 4,500 Albertans paid (nearly 1,400 children); about 5,700 evacuated",
    amount_low_cad=to25(4.6e6, 2025), amount_central_cad=to25(4.6e6, 2025), amount_high_cad=to25(4.6e6, 2025),
    method="official program total", source_name="Alberta Assisted Living and Social Services Annual Report 2025-2026, Results Analysis ('Supporting Alberta's wildfire evacuees')",
    url="https://open.alberta.ca/dataset/360bb55b-45e2-4614-8305-4fbba3e5b06e/resource/49762a21-88cd-4ec6-a06f-36156ed456c0/download/alss-annual-report-2025-2026.pdf",
    overlaps_with="DFAA (no Alberta 2025 order counted); none in gov_ or health_", additive="yes",
    notes="Payments for shelter, food, clothing and transport, not income: no overlap with lost wages. The ministry was renamed from Seniors, Community and Social Services.")
PSES = os.path.join(R2, "ab", "pses-annual-report-2025-2026.txt")
check(PSES, "the 2025 County of Grande Prairie wildfire for $1.7 million", "Town of Swan Hills and the Municipal District of Opportunity for $4.3 million",
      "Westlock County for $31.1 million")
add(item="Evacuations and response: Alberta HARP programs approved for the 2025 wildfires (high case only)", jurisdiction="AB",
    period="2025 (FY2025-26)", quantity="3 programs: County of Grande Prairie 1.7 M; Swan Hills and MD of Opportunity 4.3 M; Athabasca, Sturgeon, Bonnyville, Westlock 31.1 M",
    amount_low_cad=0, amount_central_cad=0, amount_high_cad=to25(37.1e6, 2025),
    method="approved Hazard Assistance and Resilience Program amounts (uninsurable losses and local response); not split by the province",
    source_name="Alberta Public Safety and Emergency Services Annual Report 2025-2026, pp. 36-37",
    url="https://open.alberta.ca/dataset/2b228e33-89ac-42d0-9c60-218d9ffb9f30/resource/b74e4b2b-7a28-4301-b8ae-7be28cca10f6/download/pses-annual-report-2025-2026.pdf",
    overlaps_with="property and infrastructure recovery (econ_ and building lines); DFAA", additive="yes",
    notes="Same treatment as the first-round Alberta 2023 DRP remainder: high case only.")

# Alberta 2023 Red Cross matching (rule x official count of donations eligible for matching)
PSN = os.path.join(R1, "pages", "ps_committee_note_donation_matching_20240322.html")
check(PSN, "Both federal and provincial donation matching periods were for a 30-day period, from May 8, 2023, to June 6, 2023",
      "the organization raised $7,626,270, which is eligible to be matched", "the CRC had raised $5,005,180")
check(os.path.join(R1, "ab", "pses_2023_24.txt"), "Canadian Red Cross Matching Donation")
check(os.path.join(R2, "redcross", "crc_2023_alberta_update_oct2025.txt"), "MATCHING FUNDS FROM THE GOVERNMENT OF ALBERTA: $7.4M")
add(item="Evacuations and relief: Government of Alberta match of Red Cross 2023 Alberta Fires Appeal donations", jurisdiction="AB",
    period="2023 (FY2023-24)", quantity="donations of 7,626,270 in the May 8 to June 6, 2023 matching window, matched one for one",
    amount_low_cad=to25(7.4e6, 2023), amount_central_cad=to25(7626270, 2023), amount_high_cad=to25(7626270, 2023),
    method="ESTIMATE: Alberta's commitment to match every dollar in the window (federal and provincial periods identical) x the donations Public Safety Canada counted as eligible; low = the Red Cross's own report of Alberta matching received ($7.4M)",
    source_name="Public Safety Canada Parliamentary Committee Note 'Donation Matching Efforts in Response to Domestic Emergencies' (2024-03-22); Public Safety Canada news release 2023-05-11 (Alberta and Canada each match every dollar); Alberta PSES Annual Report 2023-24 (AEMA over budget incl. 'Canadian Red Cross Matching Donation'); Canadian Red Cross 2023 Alberta Wildfires update, Oct 31 2025 (recipient report, not a government record)",
    url="https://www.publicsafety.gc.ca/cnt/trnsprnc/brfng-mtrls/prlmntry-bndrs/20240322/09-en.aspx ; https://www.canada.ca/en/public-safety-canada/news/2023/05/alberta-and-canada-to-match-canadian-red-cross-donations-to-support-alberta-wildfire-response.html ; https://cdn.redcross.ca/prodmedia/crc/azure/how-we-help/current-responses/2023%20alberta%20wildfires/2023-alberta-wildfires-update-october-2025.pdf",
    overlaps_with="federal match already in gov_ (Public Safety Canada rows), this is the provincial match; may sit inside Alberta's $175M 2023 DRP figure used in the first-round high-case remainder row (high case may then count it twice, up to 8 M)",
    additive="yes", notes="Alberta's own payment record was not found: searched Alberta grant payments disclosure 2023-24, 2024-25, 2025-26 and the 2023-24 to 2025-26 Blue Book supplier payments (no wildfire Red Cross payment listed).")
# Nova Scotia 2023 Red Cross matching and $500 household payments
NSM = os.path.join(R1, "ns", "governments-match-donations-red-cross-support-nova-scotians.html")
check(NSM, "matching every dollar donated to the Canadian Red Cross Nova Scotia and Atlantic Canada Fires Appeal")
check(os.path.join(R2, "redcross", "crc_2023_ns_atlantic_update_mar2025.txt"), "MATCHING FUNDS FROM THE GOVERNMENT OF NOVA SCOTIA: $4.6M")
add(item="Evacuations and relief: Government of Nova Scotia match of Red Cross 2023 Nova Scotia and Atlantic Canada Fires Appeal donations",
    jurisdiction="NS", period="2023 (FY2023-24)", quantity="donations of 5,005,180 in the May 31 to June 29, 2023 matching window",
    amount_low_cad=to25(4.6e6, 2023), amount_central_cad=to25(5005180, 2023), amount_high_cad=to25(5005180, 2023),
    method="ESTIMATE: Nova Scotia's commitment to match every dollar donated x donations counted by Public Safety Canada in the matching window; low = Red Cross report of Nova Scotia matching received ($4.6M)",
    source_name="Nova Scotia news release 2023-06-06 'Governments Match Donations to Red Cross to Support Nova Scotians'; Public Safety Canada committee note (2024-03-22); Canadian Red Cross Nova Scotia and Atlantic Canada Fires update, March 31 2025 (recipient report)",
    url="https://news.novascotia.ca/en/2023/06/06/governments-match-donations-red-cross-support-nova-scotians ; https://www.publicsafety.gc.ca/cnt/trnsprnc/brfng-mtrls/prlmntry-bndrs/20240322/09-en.aspx ; https://cdn.redcross.ca/prodmedia/crc/azure/how-we-help/current-responses/2023%20atlantic%20canada%20wildfire%20responses/novascotia_atlantic_wildfiresappealupdate_en.pdf",
    overlaps_with="federal match in gov_ (Public Safety Canada); DFAA NS 2023", additive="yes",
    notes="Nova Scotia's own payment record (Public Accounts 2023-24 Supplementary Information, payments to the Canadian Red Cross) could not be downloaded: notices.novascotia.ca reset every connection and the WebFetch copy was truncated.")
check(os.path.join(R1, "ns", "support-nova-scotians-impacted-wildfires-provincewide-burn-ban.html"),
      "the Province will provide $500 to every household required to evacuate", "The funding will be administered through the Canadian Red Cross")
check(os.path.join(R1, "ns", "governments-match-donations-red-cross-support-nova-scotians.html"),
      "about 20,000 Nova Scotians were evacuated from their homes and businesses in Halifax and Shelburne counties")
HRFE = os.path.join(R2, "days2", "hrm_240806rci01.txt")
check(HRFE, "required the evacuation of 16,400 people", "requiring the evacuation of nearly 7,000")
SHEL = ["2021A00051201001", "2021A00051201004", "2021A00051201006", "2021A00051201008", "2021A00051201009"]   # Shelburne County CSDs
PPH_HFX = POP("2021A00051209034") / HH["2021A00051209034"]
PPH_SHEL = sum(POP(c) for c in SHEL) / sum(HH[c] for c in SHEL)
assert abs(PPH_HFX - 2.309) < 0.001 and abs(PPH_SHEL - 2.117) < 0.001
NSH = (20000 / PPH_HFX, 16400 / PPH_HFX + 7000 / PPH_SHEL, (16400 + 7000) / PPH_SHEL)   # evacuated households, low / central / high
add(item="Evacuations: Nova Scotia $500 payment to every household required to evacuate (administered by the Red Cross)", jurisdiction="NS",
    period="2023 (FY2023-24)", quantity=f"about {NSH[1]:,.0f} evacuated households ({NSH[0]:,.0f} to {NSH[2]:,.0f})", unit_cost="500 per household",
    amount_low_cad=to25(NSH[0] * 500, 2023), amount_central_cad=to25(NSH[1] * 500, 2023), amount_high_cad=to25(NSH[2] * 500, 2023),
    method=f"ESTIMATE from official inputs, shown for reference only: the province's rule x evacuated households; households = official evacuee counts / Census 2021 people per occupied private dwelling (Halifax {PPH_HFX:.2f}, Shelburne County municipalities {PPH_SHEL:.2f}); central = 16,400 (Upper Tantallon) / {PPH_HFX:.2f} + 7,000 (Shelburne County) / {PPH_SHEL:.2f}; low = 'about 20,000' / {PPH_HFX:.2f}; high = 23,400 / {PPH_SHEL:.2f}; assumes every eligible household was paid",
    source_name="Nova Scotia news releases 2023-05-29 (rule) and 2023-06-06 (about 20,000 evacuated); Halifax Regional Fire and Emergency Post Incident Analysis (Regional Council 2024-08-06, item 240806rci01: 16,400 and nearly 7,000 evacuated); Statistics Canada Census Profile 2021 (population and occupied private dwellings by CSD)",
    url="https://news.novascotia.ca/en/2023/05/29/support-nova-scotians-impacted-wildfires-provincewide-burn-ban ; https://news.novascotia.ca/en/2023/06/06/governments-match-donations-red-cross-support-nova-scotians ; https://www.halifax.ca/sites/default/files/documents/city-hall/regional-council/240806rci01.pdf ; https://api.statcan.gc.ca/census-recensement/profile/sdmx/rest/data/STC_CP,DF_CSD/A5..1.5.1",
    overlaps_with="DFAA NS 2023; inside the NS Emergency Management Office 2023-24 disaster assistance line (gov2_, reference, not counted); not in gov_",
    additive="no",
    notes="NOT COUNTED: no official record of the amount paid exists in the documents searched, and the only reported total (3.6 M to 7,200 households by mid-June 2023, a Red Cross statement relayed by CTV News) is news, so it is not used. Searched: NS news releases 2023-05-29, 06-06, 06-22, 2023-09-28 (first budget update), 2024-07-26 (Public Accounts 2023-24) and the daily wildfire updates 2023-05-29 to 06-12; Summary of Additional Appropriations Sept 2023, Dec 2023, Feb 2024; Accountability Reports 2023-24 of Municipal Affairs and Housing (EMO: 78.3 M disaster assistance for floods and wildfires, no split) and of Labour, Skills and Immigration (no wildfire item); NS Public Accounts 2023-24 Supplementary Information (pa-supplementary-2024.pdf: notices.novascotia.ca resets shell connections and WebFetch returns only pages 1 to 28, the MAH section starts at p. 246); Red Cross Nova Scotia appeal update 2025-03-31 (appeal funds only); nslegislature.ca (connection reset; Public Accounts Committee transcript not reachable). Living-cost support, not income.")
# Manitoba 2025 Red Cross matching
MBN = os.path.join(R2, "mb", "mb_news_69482.html")
check(MBN, "both the federal government and provincial government will match every dollar donated to the Canadian Red Cross 2025 Manitoba Wildfires Appeal",
      "Donation matching will be open for 30 days, retroactive to when each appeal first opened on May 28",
      "the provincial government will match donations up to $15 million")
add(item="Evacuations and relief: Government of Manitoba match of Red Cross 2025 Manitoba Wildfires Appeal donations", jurisdiction="MB",
    period="2025 (FY2025-26)", quantity="donations in the May 28 to June 27, 2025 window: between 6.9 M (linear between $3.3M on June 6 and $8.9M on July 9) and 8.9 M",
    amount_low_cad=to25(6.86e6, 2025), amount_central_cad=to25(7.88e6, 2025), amount_high_cad=to25(8.9e6, 2025),
    method="ESTIMATE: Manitoba's commitment (match every dollar for 30 days from May 28, cap $15M) x donations raised in the window, interpolated from the Red Cross's dated totals ($3.3M as of June 6; $8.9M as of July 9; $9.0M as of July 25, all excluding matching)",
    source_name="Manitoba news release 2025-06-04 'Governments of Canada and Manitoba Match Red Cross Donations'; Canadian Red Cross media information updates of 2025-06-08, 2025-07-11 and 2025-08-01 (recipient releases, read with WebFetch)",
    url="https://news.gov.mb.ca/news/print,index.html?item=69482 ; https://www.redcross.ca/about-us/media-news/news-releases/media-information-update-manitoba-and-saskatchewan-wildfires-aa68bfc8bb9b429b630609d664d34a89 ; https://www.redcross.ca/about-us/media-news/news-releases/media-information-update-manitoba-and-saskatchewan-wildfires-july11",
    overlaps_with="likely inside Manitoba's audited $80M 2025 'evacuee related costs' (health_costs.csv, additive) or its $12M 'DFA program and other operating costs'; the federal match (up to $30M) is not yet disclosed in the federal grants database",
    additive="no", notes="Reference only, to count once: Manitoba's audited wildfire table (Public Accounts 2025-26, p. 12) covers all 2025 wildfire response spending.")
# Quebec 2023 evacuee allowances
QRA = os.path.join(R2, "qc", "RA_gestion_MSP_2023-2024.txt")
check(QRA, "une prestation d'évacuation d'urgence s'élevant à 1 500 $ par résidence évacuée", "8 911 demandes d'indemnité",
      "plus de 12 522 000 $ ont été versés en quelques semaines")
add(item="Evacuations: Quebec emergency evacuation benefit (1,500 per evacuated principal residence, each evacuation)", jurisdiction="QC",
    period="2023 (FY2023-24)", quantity="more than 8,911 claims processed", unit_cost="1,500 per residence per evacuation",
    amount_low_cad=to25(12.522e6, 2023), amount_central_cad=to25(12.522e6, 2023), amount_high_cad=to25(8911 * 1500, 2023),
    method="official amount paid ('plus de 12 522 000 $ ont été versés'); high = 8,911 claims x 1,500",
    source_name="Ministère de la Sécurité publique du Québec, Rapport annuel de gestion 2023-2024, p. 13 ('Programme spécifique et adaptation rapide des outils de réclamation lors des feux de forêt'); décret and news release 2023-06-09",
    url="https://cdn-contenu.quebec.ca/cdn-contenu/adm/min/securite-publique/publications-adm/rapport-annuel-de-gestion/RA_gestion_MSP_2023-2024.pdf ; https://www.quebec.ca/nouvelles/actualites/details/un-nouveau-programme-dassistance-financiere-pour-les-sinistres-des-incendies-de-foret-48601",
    overlaps_with="DFAA Quebec 2023 order (not added); gov_ Quebec line is SOPFEU firefighting only", additive="yes",
    notes="'More than' amounts: a floor. The same program reimbursed municipalities (100%) and organizations such as the Red Cross; those amounts are not published separately.")
# Ontario hosting of First Nations evacuees (inside ISC EMAP): official counts, cost estimate
ON25 = os.path.join(R2, "on", "on_2025_annual_report.html")
check(ON25, "The province organized the safe evacuation, hosting and return of 8,387 people affected by wildland fires.",
      "This included 6,080 residents from Ontario First Nation communities and 2,307 evacuees from Manitoba.",
      "On average, people who were evacuated because of wildland fires were hosted for approximately one month.")
check(os.path.join(R2, "on", "on_2023_annual_report.html"), "3,519 people evacuated", "more than 600 First Nation community members")
check(os.path.join(R2, "on", "on_2024_annual_report.html"), "1,316 people evacuated")
check(os.path.join(R1, "isc", "emap_key_stats.html"), "For 2024 to 2025: Response: $294,372,464", "For 2025 to 2026: Response: $688,648,240")
check(os.path.join(AN, "losses", "raw", "isc", "evac_stats.html"), "2024 to 2025: 26 short term", "2025 t0 2026: 52 short term, 21 long term*",
      "2024 to 2025: 5 short term", "2025 to 2026: 7 short term, 1 long term*")
QP = pd.read_csv(os.path.join(AN, "losses", "raw", "isc", "qpnotes_wildfire_subset.csv"))
qtxt = " ".join(QP[QP.reference_number == "ISC-2026-QP-00844"].background_en.astype(str))
assert "displacing over 44,852 people from 60 communities" in qtxt
ISC_PER = {2024: 294372464 * 26 / 31 / 15469, 2025: 688648240 * 73 / 81 / 44852}   # wildfire response $ per FN evacuee
lo_pe, hi_pe = min(ISC_PER.values()), max(ISC_PER.values())
# Ontario's own records: expense line 'First Nations Emergency Evacuations on behalf of ISC' and ISC's payments to Ontario
ON3 = os.path.join(R2, "on3")
check(os.path.join(ON3, "budget_ontario_ca_2026_chapter-3_html.html"), "Table 3.10 Total Expense 1 ($ Millions) Ministry Expense Actual 2023\u201324 Actual 2024\u201325 Interim 2025\u201326 Plan 2026\u201327",
      "Emergency Preparedness and Response \u2014 First Nations Emergency Evacuations on behalf of Indigenous Services Canada 11 0 227 \u2013")
check(os.path.join(ON3, "budget_ontario_ca_2025_fallstatement_chapter-3_html.html"), "$99.3 million to support evacuation services for First Nation communities affected by spring flooding and wildland fires")
check(os.path.join(ON3, "on_2025-26_third_quarter.html"), "$50.7 million to support evacuation services for First Nation communities affected by spring flooding and wildland fires")
check(os.path.join(ON3, "pa_ms_2025-26.txt"), "First Nations Emergency Assistance Program................................................. 214,949,965 13,308,513",
      "TOTAL OPERATING EXPENSE FOR EMERGENCY MANAGEMENT 63,180,214 237,766,300 300,946,514 ONTARIO PROGRAM........................ 270,876,874")
rev = []
for fy in ("2023-24", "2024-25", "2025-26"):           # column name differs by year ('Amount $', ' Amount')
    d = pd.read_csv(os.path.join(ON3, f"pa_ms_revenue_{fy}.csv"), encoding="utf-8-sig")
    rev.append(d.rename(columns=lambda c: "Amount" if c.strip().startswith("Amount") else c.strip()))
rev = pd.concat(rev)
rev = rev[rev["Revenue Detail"].eq("First Nations Emergency Assistance Program")].set_index("Year").Amount.astype(float)
assert rev.to_dict() == {"2023-24": 93432, "2024-25": 13308513, "2025-26": 214949965}
check(os.path.join(R2, "on", "on_2025_annual_report.html"), "Ontario also co-ordinated the evacuation of 11,069 people from 10 communities in response to a wide range of hazards.")
ONSH = {2023: (600, 3519 - 600), 2025: (8387, 11069 - 8387)}     # wildfire evacuees, other-hazard evacuees (Ontario annual reports)


def fire_share(w, o):
    """wildfire share of evacuation cost: people share (central); other-hazard evacuees hosted twice (low) or half (high) as long"""
    return w / (w + 2 * o), w / (w + o), w / (w + 0.5 * o)


s23, s25 = fire_share(*ONSH[2023]), fire_share(*ONSH[2025])
add(item="Evacuations: Ontario First Nations emergency evacuations on behalf of ISC, wildfire share (inside ISC EMAP)", jurisdiction="ON",
    period="2023 (FY2023-24)", quantity="Ontario expense 11 M, all hazards (3,519 evacuees in 2023, of whom more than 600 First Nation members for wildland fires)",
    unit_cost=f"wildfire share {s23[0]:.2f} / {s23[1]:.2f} / {s23[2]:.2f}",
    amount_low_cad=to25(11e6 * s23[0], 2023), amount_central_cad=to25(11e6 * s23[1], 2023), amount_high_cad=to25(11e6 * s23[2], 2023),
    method="official actual expense (all hazards, rounded to $ millions) x ESTIMATED wildfire share: wildfire evacuees / all evacuees (central); low and high assume the other-hazard (flood) evacuees were hosted twice or half as long",
    source_name="Ontario 2026 Budget, Chapter 3, Table 3.10 (actual 2023-24); Ontario Public Accounts ministry statements, revenue (ISC paid 93,432 in FY2023-24 and 13,308,513 in FY2024-25); Safe, Practiced and Prepared Ontario 2023 annual report",
    url="https://budget.ontario.ca/2026/chapter-3.html ; https://data.ontario.ca/dataset/df804071-1a37-4eff-a96c-7e267ac02cc9 ; https://www.ontario.ca/page/safe-practiced-and-prepared-ontario-2023-annual-report",
    overlaps_with="INSIDE the ISC EMAP rows in gov_costs.csv and gov2_costs.csv (ISC reimburses Ontario 100%): count once there",
    additive="no", notes="Replaces the first version of this row (per-evacuee estimate). In 2023-24 Emergency Management Ontario sat in Treasury Board Secretariat; the line is restated under MEPR in the 2026 Budget.")
add(item="Evacuations: Ontario First Nations emergency evacuations on behalf of ISC, wildfire share (inside ISC EMAP)", jurisdiction="ON",
    period="2024 (FY2024-25)", quantity="Ontario expense 0 M (rounded); Ontario reported no wildfire evacuations it hosted in 2024 (1,316 evacuated, all hazards)",
    amount_low_cad=0, amount_central_cad=0, amount_high_cad=to25(0.5e6, 2024),
    method="official actual expense, rounded to $ millions (so below 0.5 M); high = the rounding limit",
    source_name="Ontario 2026 Budget, Chapter 3, Table 3.10 (actual 2024-25); Safe, Practiced and Prepared Ontario 2024 annual report",
    url="https://budget.ontario.ca/2026/chapter-3.html ; https://www.ontario.ca/page/safe-practiced-and-prepared-ontario-2024-annual-report",
    overlaps_with="INSIDE the ISC EMAP rows (gov_, gov2_)", additive="no",
    notes="ISC's 13.3 M payment to Ontario in FY2024-25 most likely settles 2023-24 costs (Ontario booked 11 M of expense in 2023-24 and received 0.09 M that year).")
on25 = (270876874 - 63180214, float(rev["2025-26"]), 227e6)      # expense above the EMO estimate; ISC payment; interim expense
add(item="Evacuations: Ontario First Nations emergency evacuations on behalf of ISC, wildfire share (inside ISC EMAP)", jurisdiction="ON",
    period="2025 (FY2025-26)", quantity=f"ISC paid Ontario {on25[1] / 1e6:.1f} M (FY2025-26 actual); 11,069 evacuees from 10 communities, 8,387 for wildland fires (6,080 Ontario First Nations, 2,307 from Manitoba), hosted about one month",
    unit_cost=f"wildfire share {s25[0]:.2f} / {s25[1]:.2f} / {s25[2]:.2f}",
    amount_low_cad=to25(on25[0] * s25[0], 2025), amount_central_cad=to25(on25[1] * s25[1], 2025), amount_high_cad=to25(on25[2] * s25[2], 2025),
    method="all-hazard cost x ESTIMATED wildfire share (as for 2023); cost: central = ISC payments received in FY2025-26 (Public Accounts, actual); low = Emergency Management Ontario operating expense above its printed estimate (270.9 M - 63.2 M); high = the 2026 Budget interim for the line (227 M)",
    source_name="Public Accounts of Ontario 2025-26, Ministry of Emergency Preparedness and Response, Statement of Revenue (p. 2-131) and Statement of Expenses (p. 2-127); Ontario 2026 Budget, Table 3.10; 2025 Fall Economic Statement (99.3 M added) and 2025-26 Third Quarter Finances (50.7 M added); Safe, Practiced and Prepared Ontario 2025 annual report",
    url="https://www.ontario.ca/files/2026-10/tbs-public-ministry-statements-and-schedules-2025-2026-en-2006-10-13.pdf ; https://budget.ontario.ca/2026/chapter-3.html ; https://budget.ontario.ca/2025/fallstatement/chapter-3.html ; https://www.ontario.ca/page/2025-26-third-quarter-finances ; https://www.ontario.ca/page/safe-practiced-and-prepared-ontario-2025-annual-report",
    overlaps_with="INSIDE the ISC EMAP rows in gov_costs.csv and gov2_costs.csv (ISC reimburses Ontario 100%): count once there",
    additive="no", notes=f"Replaces the first version of this row (per-evacuee estimate, 116 to 134 M). All-hazard total: ISC paid {on25[1] / 1e6:.1f} M. Sandy Lake First Nation's June 2025 evacuation is in the 8,387. The 2026-27 line (140 M at Q1, mostly spring floods) is a separate row below.")
# NWT final ETSP and EIDS (audited actuals), other GNWT evacuation items, and the GNWT 204.6 M split (partial)
TD31 = os.path.join(R2, "nwt", "td_31-201.txt")
check(TD31, "EVACUATION SUPPORT PROGRAM DATA as of February 14, 2024", "Total Expenditures 3,612,000 6,061,850",
      "Processed for Payment 4,817 8,712")
TD29 = os.path.join(R2, "nwt", "td_29-201.txt")
check(TD29, "NWT Evacuation Travel Support Program 6,000,000", "2022 Flood and 2023 Wildfire Response and Recovery 55,153,000",
      "Operational Support for Businesses Affected by Wildfire Evacuation 1,000,000 (1,000,000)", "2023-24 Supplemental Fire Suppression Funding 24,300,000")
TD967 = os.path.join(R1, "nwt", "supp_no2.txt")
check(TD967, "TD 967-19(2) TABLED ON AUGUST 28, 2023", "Fire Suppression Shortfall for 2023 Wildfire season 75,000,000")
TD974 = os.path.join(R2, "nwt", "td_974-192.txt")
check(TD974, "TD 974-19(2) TABLED ON SEPTEMBER 27, 2023", "To increase the authorized short term borrowing limit of the Government to accommodate unplanned",
      "expenditures as a result of fire related emergencies.", "Total Government Borrowing Authorization 1,259,345,000 120,000,000 1,379,345,000")
PANT = os.path.join(RG, "nt", "2023-2024_section_ii_-_non-consolidated_financial_statements.txt")
check(PANT, "Non-Consolidated Schedule of Grants (unaudited) Schedule 5", "Evacuation Income Support Program - - - - 3,684 (3,684)",
      "Evacuation Travel Support Program - 6,000 - 6,000 6,110 (110)", "United Way NWT - Fire Relief - - - - 400 (400)",
      "There were no Special Warrants for the fiscal year ended March 31, 2024.")
ME26 = os.path.join(RG, "nt", "me2526.txt")
check(ME26, "Percentage 2024-2025 2024-2025 2025-2026 Externally 2023-2024 Main Revised Main Funded Actuals Estimates Estimates Estimates",
      "Support for Entrepreneur and Economic Development - Operation Support for Businesses Affected by Wildfire Evacuation 50% 2,279",
      "2023 Wildfire 91,645 - - -")
ETSP, EIDS, UWAY, ITIB = 6.110e6, 3.684e6, 0.400e6, 2.279e6      # FY2023-24 actuals ($ thousands in the source)
PA_URL = "https://www.fin.gov.nt.ca/sites/fin/files/resources/2023-2024_section_ii_-_non-consolidated_financial_statements.pdf"
ME_URL = "https://www.ntlegislativeassembly.ca/sites/default/files/tabled-documents/2025-02/TD%20275-20%281%29%202025-2026%20Main%20Estimates.pdf"
old_etsp = H1[H1.item.str.startswith("Evacuations: NWT Evacuation Travel Support Program")].iloc[0]
add(item="Evacuations: NWT Evacuation Travel Support Program, final", jurisdiction="NT", period="2023 (FY2023-24)",
    quantity="8,712 applications processed for payment (9,515 received) by 2024-02-14", unit_cost="750 per vehicle leaving the NWT, 400 within",
    amount_low_cad=to25(ETSP, 2023), amount_central_cad=to25(ETSP, 2023), amount_high_cad=to25(ETSP, 2023),
    method="official FY2023-24 actual expenditure (grant schedule of the Public Accounts); 6,061,850 at 2024-02-14 (TD 31-20(1))",
    source_name="GNWT Public Accounts 2023-24, Section II, Schedule 5 Non-Consolidated Schedule of Grants (Finance: 6,110 thousand); Tabled Document 31-20(1) (counts); TD 29-20(1) Supplementary Estimates (Operations) No. 4, 2023-24 (6.0 M appropriated)",
    url=PA_URL + " ; https://www.ntlegislativeassembly.ca/sites/default/files/legacy/td_31-201.pdf ; https://www.ntlegislativeassembly.ca/sites/default/files/legacy/td_29-201.pdf",
    overlaps_with=f"REPLACES health_costs.csv row 'Evacuations: NWT Evacuation Travel Support Program (partial, to 2023-09-05)' (central {float(old_etsp.amount_central_cad) / 1e6:.2f} M): drop it when adding this; inside GNWT's $204.6M estimate (see the high-case adjustment row); DFAA",
    additive="yes", notes="Audited-year actual; no special warrants were needed in FY2023-24 (Schedule 7).")
add(item="Income support: NWT Evacuee Income Disruption Support, final", jurisdiction="NT", period="2023 (FY2023-24)",
    quantity="4,817 applications processed (6,336 received) by 2024-02-14", unit_cost="750 one-time",
    amount_low_cad=to25(EIDS, 2023), amount_central_cad=to25(EIDS, 2023), amount_high_cad=to25(EIDS, 2023),
    method="official FY2023-24 actual expenditure (grant schedule); 3,612,000 at 2024-02-14 (TD 31-20(1))",
    source_name="GNWT Public Accounts 2023-24, Section II, Schedule 5 (Finance: Evacuation Income Support Program 3,684 thousand); TD 31-20(1)",
    url=PA_URL + " ; https://www.ntlegislativeassembly.ca/sites/default/files/legacy/td_31-201.pdf",
    overlaps_with="lost work and wages (income replacement); replaces the first-round partial EIDS reference row", additive="no",
    notes="Transfer that replaces lost income, which is counted in lost wages.")
add(item="Evacuations and relief: GNWT grant to United Way NWT for fire relief", jurisdiction="NT", period="2023 (FY2023-24)",
    amount_low_cad=to25(UWAY, 2023), amount_central_cad=to25(UWAY, 2023), amount_high_cad=0,
    method="official FY2023-24 actual grant (unbudgeted)", source_name="GNWT Public Accounts 2023-24, Section II, Schedule 5 (Finance: United Way NWT - Fire Relief 400 thousand)",
    url=PA_URL, overlaps_with="inside health_costs.csv 'GNWT wildfire emergency costs not itemized above (high case only)', so counted in low and central only; the federal match of United Way NWT donations is a separate federal payment (gov_ Public Safety Canada row)",
    additive="partial", notes="Same treatment as gov2_ Yellowknife and Hay River rows: additive in low and central, dropped in the high case.")
PANT25 = os.path.join(RG, "nt", "2024-2025_public_accounts_section_ii_-_non-consolidated_financial_statements.txt")
check(PANT25, "Emergency Evacuation Relief Program - - - - 212 (212)")
check(os.path.join(R1, "nwt", "emar2025.txt"), "These included evacuations in Fort Good Hope (June\u2013July 2024), Whatì (August\u2013September 2025), and Fort Providence (August\u2013September 2025).")
add(item="Evacuations: NWT Emergency Evacuation Relief Program, FY2024-25 (Fort Good Hope 2024 evacuation)", jurisdiction="NT", period="2024 (FY2024-25)",
    quantity="Fort Good Hope, about 380 residents evacuated June 15 to July 6, 2024 (the only NWT evacuation that fiscal year)",
    amount_low_cad=to25(0.212e6, 2024), amount_central_cad=to25(0.212e6, 2024), amount_high_cad=0,
    method="official FY2024-25 actual grant (unbudgeted), Finance", source_name="GNWT Public Accounts 2024-25, Section II, Schedule 5 Non-Consolidated Schedule of Grants; NWT Emergency Management Annual Report 2025 (evacuations April 2024 to 2025)",
    url="https://www.fin.gov.nt.ca/sites/fin/files/resources/2024-2025_public_accounts_section_ii_-_non-consolidated_financial_statements.pdf ; https://www.maca.gov.nt.ca/sites/maca/files/resources/nwt_emergency_management_annual_report_2025.pdf",
    overlaps_with="inside health_costs.csv 'GNWT wildfire emergency costs not itemized above (high case only)' 2024 (the 7.4 M estimate), so counted in low and central only",
    additive="partial", notes="Living-cost relief for evacuees, like the 2023 travel support; the 2025 payments (Whati, Fort Providence) will appear in the 2025-26 Public Accounts, not yet tabled.")
add(item="Business support: NWT operational support for businesses affected by the wildfire evacuation (reference)", jurisdiction="NT", period="2023 (FY2023-24)",
    amount_low_cad=to25(ITIB, 2023), amount_central_cad=to25(ITIB, 2023), amount_high_cad=to25(ITIB, 2023),
    method="official FY2023-24 actual expenditure under a one-time agreement with CanNor (50% externally funded)",
    source_name="GNWT 2025-2026 Main Estimates (TD 275-20(1)), Industry, Tourism and Investment, Cost Shared Agreements (information item), 2023-2024 actuals; TD 29-20(1) (1.0 M appropriated, offset by federal revenue)",
    url=ME_URL + " ; https://www.ntlegislativeassembly.ca/sites/default/files/legacy/td_29-201.pdf",
    overlaps_with="business losses in econ_ (GDP and insured business interruption): a transfer to businesses, like the Jasper and Quebec business supports that econ_ marks 'no (transfer)'; CanNor's half is federal money",
    additive="no", notes="Transfer that compensates business losses; not a cost of its own.")
add(item="GNWT 2023-24 wildfire appropriations and federal revenue (reference)", jurisdiction="NT", period="2023 (FY2023-24)",
    quantity="Supplementary Estimates No. 2 (TD 967-19(2), 2023-08-28): fire suppression shortfall 75.0 M; No. 3 (TD 974-19(2), 2023-09-27): borrowing authority +120 M for 'unplanned expenditures as a result of fire related emergencies', no spending items; No. 4 (TD 29-20(1), 2024-02-20): fire suppression 24.3 M, ETSP 6.0 M, MACA 2022 flood and 2023 wildfire response and recovery 55.153 M, ITI business support 1.0 M; MACA revenue 'Other Federal Transfers: 2023 Wildfire' 91.645 M (2023-24 actual)",
    dollars_basis="", method="reference: official appropriation documents and revenue summary",
    source_name="NWT Tabled Documents 967-19(2), 974-19(2), 29-20(1); GNWT 2025-2026 Main Estimates, Municipal and Community Affairs Revenue Summary; GNWT Public Accounts 2023-24 Schedule 7 (no special warrants)",
    url="https://www.ntlegislativeassembly.ca/sites/default/files/legacy/td_974-192.pdf ; https://www.ntlegislativeassembly.ca/sites/default/files/legacy/td_29-201.pdf ; " + ME_URL + " ; " + PA_URL,
    overlaps_with="fire suppression is in gov_ (ECC Forest Management actual 150.2 M); the 91.6 M federal transfer reimburses GNWT costs (most likely DFAA payments), so it is not a separate cost", additive="no",
    notes="Supplementary Estimates No. 3 holds no expenditure split: it only raised short-term borrowing. No GNWT document found splits the $204.6M by purpose (see the next row).")
check(os.path.join(R1, "ab", "pses_2023_24.txt"), "$41.6 million increase in Other Revenue from 2023 NWT wildfires evacuation cost recovery")
check(os.path.join(R1, "nwt", "emar2025.txt"), "2023 Wildfires with an estimated $204.6 million in costs",
      "Confirming Alberta")
G2 = pd.read_csv(os.path.join(CF, "gov2_costs.csv"))          # other agent's second round, read only
yk = float(G2[G2.item.str.startswith("City of Yellowknife 2023 wildfire response costs")].amount_nominal_cad.iloc[0])
hr = float(G2[G2.item.str.startswith("Town of Hay River 2023 wildfire civil emergency costs")].amount_nominal_cad.iloc[0])
ident = 41.6e6 + ETSP + EIDS + UWAY + ITIB + yk + hr
add(item="GNWT 2023 wildfire event cost ($204.6M): components identified in official records (partial split, reference)", jurisdiction="NT", period="2023",
    quantity=f"Alberta hosting recovered from the NWT 41.6 M (Alberta FY2023-24 revenue) + ETSP 6.11 M + EIDS 3.68 M + United Way grant 0.40 M + ITI business support 2.28 M + community claims Yellowknife {yk / 1e6:.2f} M and Hay River {hr / 1e6:.2f} M = {ident / 1e6:.1f} M ({100 * ident / 204.6e6:.0f}%)",
    amount_low_cad=to25(ident, 2023), amount_central_cad=to25(ident, 2023), amount_high_cad=to25(ident, 2023),
    method=f"sum of components reported separately by GNWT, Alberta and the two municipalities; the remainder ({(204.6e6 - ident) / 1e6:.0f} M) is other community claims, recovery projects and GNWT department costs that no published document splits; whether ITI business support sits inside the 204.6 M is not stated",
    source_name="NWT Emergency Management Annual Report 2025, Table 7; GNWT Public Accounts 2023-24 Schedule 5; GNWT 2025-26 Main Estimates (ITI); Alberta PSES Annual Report 2023-24; City of Yellowknife and Town of Hay River audited statements (as in gov2_costs.csv)",
    url="https://www.maca.gov.nt.ca/sites/maca/files/resources/nwt_emergency_management_annual_report_2025.pdf ; " + PA_URL + " ; https://open.alberta.ca/dataset/2b228e33-89ac-42d0-9c60-218d9ffb9f30/resource/15e58606-1375-4070-97e2-b331fc44cea2/download/pses-annual-report-2023-2024.pdf",
    overlaps_with="every component is counted in its own row (health_ Alberta hosting 44.3 M, the rows above, gov2_ municipal rows)", additive="no",
    notes="Reference only. Searched for a split by purpose: NWT tabled documents 2023-24 and 2024-25 (Supplementary Estimates No. 2, 3, 4; TD 31-20(1); TD 1007-19(2)), Public Accounts 2023-24 and 2024-25 Section II, Main Estimates 2024-25 to 2026-27, NWT 2023 Wildfire Response Review (MNP, July 2024), EM Annual Report 2025; none gives one.")
adj = -(ETSP + EIDS - 3.1e6)
add(item="Evacuations and response: high-case adjustment to the first-round GNWT 2023 remainder (final ETSP and EIDS larger than the partial 3.1 M)",
    jurisdiction="NT", period="2023", amount_low_cad=0, amount_central_cad=0, amount_high_cad=to25(adj, 2023),
    method=f"first-round high-case remainder subtracted 3.1 M of ETSP and EIDS; the audited actuals are {(ETSP + EIDS) / 1e6:.2f} M, so the remainder falls by {-adj / 1e6:.2f} M",
    source_name="GNWT Public Accounts 2023-24 Schedule 5; health_costs.csv GNWT remainder row", url=PA_URL,
    overlaps_with="adjusts health_costs.csv 'GNWT wildfire emergency costs not itemized above (high case only)' 2023", additive="yes",
    notes="Negative in the high case only.")

# =================================================================== GAP 1: 2026 EVACUATIONS
check(os.path.join(R1, "bc2026", "2026EMCR0037-000977.html"), "At the time, more than 20,000 people were evacuated from their homes.")
check(os.path.join(R2, "isc", "isc_wildfire_evac_info_2026.html.txt"), "8,805 First Nations on reserve and eligible under EMAP evacuated",
      "as of October 2, 2026")
N26 = os.path.join(R2, "nwt2026")
check(os.path.join(N26, "nwt-emergency-response-update-july-18-2026.html"),
      "The Evacuation Order in place since June 28, 2026 for Fort Simpson was downgraded to an Evacuation Alert on July 18, 2026 at 8:00 a.m.",
      "The Evacuation Order was given by the community of Wrigley on June 29, 2026.")
check(os.path.join(N26, "nwt-emergency-response-update-july-24-2026.html"), "The Evacuation Order for Wrigley remains in place.")
check(os.path.join(R1, "bc2026", "rdos_5841.html"), "Three Evacuation Orders have been downgraded to Evacuation Alerts today, Sunday, August 16, 2026.")
FS, WR, FH = POP("2021A00056104038"), POP("2021A00056104044"), POP("2021A00053560053")
assert (FS, WR, FH) == (1100, 117, 977)
n26 = (20000 + FS + WR + FH, 20000 + 8805 + FS + WR, 20000 + 8805 + FS + WR)
add(item="People evacuated by wildfire (count, no dollars; PARTIAL, revised)", jurisdiction="Canada", period="2026",
    quantity=f"at least {n26[0]:,} (BC more than 20,000 at Aug 8 + NWT Fort Simpson {FS:,} and Wrigley {WR} + Eabametoong {FH}); central {n26[1]:,} (BC 20,000 + ISC 8,805 First Nations on reserve + NWT {FS + WR:,})",
    dollars_basis="", method="sum of official counts; BC's figure is a point-in-time count (people out at once on Aug 8), so the season total is higher; ISC and BC overlap for BC First Nations (share not published); populations are 2021 Census",
    source_name="BC EMCR information bulletin 2026-08-21 (2026EMCR0037); ISC 'Wildfire evacuation information for Indigenous peoples' (modified 2026-10-05); GNWT NWT Emergency Response Updates 2026-07-18 and 07-24; Public Safety Canada 2026-08-18 update (Eabametoong evacuated by the CAF); StatCan Census 2021 (CSD 6104038, 6104044, 3560053)",
    url="https://news.gov.bc.ca/34192 ; https://isc.gc.ca/eng/1686078907307/1686078925161 ; https://www.gov.nt.ca/en/newsroom/nwt-emergency-response-update-july-18-2026 ; https://www.canada.ca/en/public-safety-canada/news/2026/08/the-government-of-canada-provides-an-update-regarding-the-2026-wildfire-season--august-update.html ; https://api.statcan.gc.ca/census-recensement/profile/sdmx/rest/data/STC_CP,DF_CSD/A5..1.1.1",
    additive="no", notes="No national 2026 total is published (CWFEDB summary ends at 2025). Days away found: Fort Simpson 20 (June 28 to July 18); Wrigley at least 25 (June 29, still ordered July 24); Summerland ordered Aug 8, last listed downgrades Aug 16 (8 days); Ontario 2026 figures not yet published.")
# official 2026 cost records (reference)
check(os.path.join(RG, "on", "on_2026-27_q1.txt"), "First Nations Emergency Evacuations on behalf of Indigenous Service Canada | – | 140.0 | 140.0",
      "federal funding to provide evacuation support services to First Nation communities affected by spring flooding and other emergencies")
add(item="Evacuations: Ontario 'First Nations Emergency Evacuations on behalf of Indigenous Services Canada', 2026-27 (reference)", jurisdiction="ON",
    period="2026 (FY2026-27, outlook at June 30)", amount_low_cad=to25(140e6, 2026), amount_central_cad=to25(140e6, 2026), amount_high_cad=to25(140e6, 2026),
    method="official in-year outlook (Q1); fully funded by ISC", source_name="Ontario 2026-27 First Quarter Finances, Tables 2 and 3 and program expense update",
    url="https://www.ontario.ca/page/2026-27-first-quarter-finances", overlaps_with="ISC EMAP FY2026-27 (estimate row below); mostly spring floods",
    additive="no", notes="Mostly spring-flood evacuations (information to June 30, before the July 2026 northwestern Ontario fire evacuations); wildfire share not stated.")
check(os.path.join(RG, "mb", "mb_q1_2026-27.txt"), "Emergency Expenditures is currently projected at $277 million over budget")
check(os.path.join(RG, "bc", "bc_2026-27_q1.txt"), "Emergency and Disaster Management Act")
add(item="Evacuations 2026: other official records checked (no wildfire evacuation dollar yet)", jurisdiction="BC, MB, SK, AB", period="2026 (FY2026-27)",
    quantity="BC Q1: EDMA statutory spending forecast 36 M, unchanged from budget (information to June 30); MB Q1: Emergency Expenditures 277 M over budget, mostly June rains and floods, includes the 2026 wildfire response (no split); SK Q1 (Aug 27): no wildfire variance; AB Q1: 104.7 M wildfire presuppression and response (Forestry and Parks, firefighting)",
    dollars_basis="", method="reference", source_name="BC 2026/27 First Quarterly Report; Manitoba 2026-27 First Quarter Report; Saskatchewan 2026-27 Q1; Alberta 2026-27 Q1 Fiscal Update",
    url="raw/gov/bc/bc_2026-27_q1.pdf ; raw/gov/mb/mb_q1_2026-27.pdf ; raw/gov/sk/sk_129712.pdf ; raw/gov/ab/ab_2026-27_q1.pdf", additive="no",
    notes="BC's August 2026 Okanagan evacuations fall after the Q1 cut-off; the Q2 reports (late November) are not out.")
# 2026 estimates
add(item="Evacuations and response on reserve 2026: ISC wildfire response cost, estimate", jurisdiction="First Nations on reserve (all provinces)",
    period="2026 (FY2026-27)", quantity="8,805 First Nations evacuees (Apr 1 to Oct 2, 2026)", unit_cost=f"{lo_pe:,.0f} to {hi_pe:,.0f} per evacuee",
    amount_low_cad=0, amount_central_cad=to25(8805 * (lo_pe + hi_pe) / 2, 2026), amount_high_cad=to25(8805 * hi_pe, 2026),
    method="ESTIMATE: ISC's 2026 evacuee count x ISC wildfire response reimbursement per First Nations evacuee in FY2024-25 and FY2025-26; low = 0 (like the gov_ ISC tier C estimate, kept out of the floor)",
    source_name="ISC 'Wildfire evacuation information for Indigenous peoples' (2026-10-05); ISC EMAP key statistics; ISC evacuation statistics; ISC QP note ISC-2026-QP-00844",
    url="https://isc.gc.ca/eng/1686078907307/1686078925161 ; https://www.sac-isc.gc.ca/eng/1685047435261/1685047540188",
    overlaps_with="not in gov_ (its ISC tier C covers FY2023-24 to FY2025-26); overlaps Ontario's 2026-27 ISC-funded evacuation line and any provincial host costs ISC reimburses",
    additive="partial", notes=f"Central {8805 * (lo_pe + hi_pe) / 2 / 1e6:.0f} M nominal. FY2026-27 EMAP costs are not published until mid-2027.")
check(os.path.join(R1, "pages", "bc_news_2024EMCR0001-000004.html"), "24,300 households")
bc23 = float(H1[(H1.jurisdiction == "BC") & H1.period.eq("2023 fires")].amount_central_cad.iloc[0])
per_bc = 118.607e6 / (24300 * 2.4)          # nominal EMCR payments for 2023 fires per person under order (2.4 persons per household)
per_ab = 4.6e6 / 5700
per_mb = 80e6 / 33000
off26 = 20000 + FS + WR
add(item="Evacuations 2026 off reserve (BC, NWT): evacuee support and local response, estimate", jurisdiction="BC, NT", period="2026",
    quantity=f"{off26:,} evacuees (BC more than 20,000 + Fort Simpson + Wrigley)",
    unit_cost=f"per evacuee: AB 2025 payments {per_ab:,.0f} (low); BC 2023 EMCR payments per person under order {per_bc:,.0f} (central); MB 2025 evacuee costs {per_mb:,.0f} (high)",
    amount_low_cad=to25(off26 * per_ab, 2026), amount_central_cad=to25(off26 * per_bc, 2026), amount_high_cad=to25(off26 * per_mb, 2026),
    method="ESTIMATE: official 2026 evacuee counts x official cost per evacuee from earlier seasons: Alberta 2025 Emergency Evacuation Payments 4.6 M / 5,700; BC EMCR EDMA payments for the 2023 interface fires 118.6 M nominal / (24,300 households x 2.4 people); Manitoba 2025 evacuee related costs 80 M / 33,000",
    source_name="as cited in the rows for those programs (ALSS 2025-26; BC EMCR EDMA payment reports and 2024-01-04 release; Manitoba Public Accounts 2025-26)",
    url="see health_costs.csv BC and Manitoba rows; ALSS 2025-26 annual report", overlaps_with="DFAA; ISC (BC First Nations evacuees inside the BC count)",
    additive="yes", notes="BC's count is a peak, so the cost may be understated; 2.4 people per household is the BC 2021 Census average (assumption stated). Published costs for 2026 will replace this.")

# =================================================================== GAP 3: DAYS AWAY, LOST WORK, EI
TD1007 = os.path.join(R2, "nwt", "td_1007-192.txt")
check(TD1007, "KFN issues an alert followed by an evacuation order for full evacuation of reserve", "May 15 o Hay River issues evacuation order",
      "May 25 o Hay River evacuation order downgraded", "June 6 o KFN evacuation order rescinded", "May 31: o Sambaa K'e issues evacuation order",
      "June 30 – Sambaa K'e residents return home", "July 4 and 5", "August 3 o Evacuation order for Behchoko",
      "August 12 o Fort Smith issues an evacuation order", "September 17 – Hay River re-entry of general public",
      "September 18 – Fort Smith re-entry of general public", "September 22 – Enterprise re-entry of residents with undamaged homes",
      "August 24 o Evacuation order for Jean Marie River is lifted", "GNWT issues evacuation order for Yellowknife, N'dilo, Dettah and",
      "Evacuation order for Yellowknife, N'dilo, Dettah and Ingraham Trail downgraded to an evacuation alert",
      "September 9 – Kakisa downgrades evacuation order")
ev = [  # year, place, prov, CSD code(s), start, end, source
    (2023, "Hay River (May)", "NT", ["2021A00056105016"], "2023-05-15", "2023-05-25", "TD 1007-19(2)"),
    (2023, "Kátł'odeeche First Nation (May)", "NT", ["2021A00056104017"], "2023-05-14", "2023-06-06", "TD 1007-19(2)"),
    (2023, "Sambaa K'e", "NT", ["2021A00056104006"], "2023-05-31", "2023-06-30", "TD 1007-19(2)"),
    (2023, "Wekweètì", "NT", ["2021A00056103052"], "2023-06-29", "2023-07-04", "TD 1007-19(2) (return July 4 and 5)"),
    (2023, "Behchokǫ̀", "NT", ["2021A00056103031"], "2023-07-24", "2023-08-03", "TD 1007-19(2)"),
    (2023, "Fort Smith", "NT", ["2021A00056105001"], "2023-08-12", "2023-09-18", "TD 1007-19(2)"),
    (2023, "Hay River (August)", "NT", ["2021A00056105016"], "2023-08-13", "2023-09-17", "TD 1007-19(2)"),
    (2023, "Kátł'odeeche First Nation (August)", "NT", ["2021A00056104017"], "2023-08-13", "2023-09-17", "TD 1007-19(2) (own hosting; Hay River re-entry date assumed)"),
    (2023, "Enterprise", "NT", ["2021A00056105003"], "2023-08-13", "2023-09-22", "TD 1007-19(2)"),
    (2023, "Jean Marie River", "NT", ["2021A00056104013"], "2023-08-13", "2023-08-24", "TD 1007-19(2)"),
    (2023, "Yellowknife, N'dilo, Dettah", "NT", ["2021A00056106023", "2021A00056106021"], "2023-08-16", "2023-09-06", "TD 1007-19(2)"),
    (2023, "Kakisa", "NT", ["2021A00056104005"], "2023-08-17", "2023-09-09", "TD 1007-19(2)"),
]
evr = []
for y, place, p, codes, s, e, src in ev:
    n = sum(POP(c) for c in codes)
    d = (pd.Timestamp(e) - pd.Timestamp(s)).days
    evr.append(dict(year=y, place=place, prov=p, people=n, start=s, end=e, days_low=d, days=d, days_high=d, source=src))
# 2023 outside the NWT: Alberta Emergency Alerts (official CAP archive), Nova Scotia releases and HRM's post incident analysis, BC EMCR counts
AEA = os.path.join(R2, "aea")
check(os.path.join(AEA, "aea_2023-05-04T23_09_52-06_00_AEMA.html"), "Updated May 04, 2023 at 23:09 HRS", "This includes the entire Town of Drayton Valley.")
check(os.path.join(AEA, "aea_2023-05-16T09_45_22-06_00_AEMA.html"), "Brazeau County evacuation order remains in effect.", "This includes the entire community of Drayton Valley.")
check(os.path.join(AEA, "aea_2023-05-16T13_54_37-06_00_BrazeauCounty.html"), "Issued May 16, 2023 at 13:54 HRS",
      "The order has been lifted for the remainder of Brazeau County and the Town of Drayton Valley.")
check(os.path.join(AEA, "aea_2023-05-06T13_04_33-06_00_AEMA.html"), "Alert: May 06, 2023 at 08:33 HRS",
      "There is a wildfire near Fox Creek that is forcing the evacuation of the entire community and surround areas.")
check(os.path.join(AEA, "aea_20230524_fox_creek.html"), "Issued May 24, 2023 at 11:55 HRS", "The Town of Fox Creek has lifted the Evacuation Order due to Wildfire.")
check(os.path.join(AEA, "aea_2023-05-05T17_37_18-06_00_yellowhead.html"), "Issued May 05, 2023 at 17:37 HRS",
      "Yellowhead County and Town of Edson have issued an Immediate Evacuation Order for the Town of Edson and Yellowhead County.")
check(os.path.join(AEA, "aea_2023-05-08T08_16_53-06_00_yellowhead.html"), "Updated May 08, 2023 at 08:16 HRS",
      "Some areas of Yellowhead County and Town of Edson have had the evacuation order lifted.",
      "Areas still under evacuation order are east of Range Road 110 (north of Chip Lake)")
check(os.path.join(AEA, "aea_2023-06-14T18_02_41-06_00_yellowhead.html"), "Alert: Jun 09, 2023 at 18:22 HRS", "The Town of Edson remains under an Evacuation Order.")
check(os.path.join(AEA, "aea_2023-06-15T18_05_14-06_00_yellowhead.html"), "Updated Jun 15, 2023 at 18:05 HRS",
      "Yellowhead County downgraded to a 4-hour evacuation advisory for parts of Yellowhead County due to ongoing wildfires.")
for place, code, s, e, src in (("Drayton Valley", "2021A00054811031", "2023-05-04", "2023-05-16", "Alberta Emergency Alerts (Brazeau County) 2023-05-04 23:09, 05-16 09:45 and 13:54"),
                               ("Fox Creek", "2021A00054818002", "2023-05-06", "2023-05-24", "Alberta Emergency Alerts (MD of Greenview) 2023-05-06 08:33 and 2023-05-24 11:55"),
                               ("Edson (May)", "2021A00054814024", "2023-05-05", "2023-05-08", "Alberta Emergency Alerts (Yellowhead County) 2023-05-05 17:37 and 05-08 08:16"),
                               ("Edson (June)", "2021A00054814024", "2023-06-09", "2023-06-15", "Alberta Emergency Alerts (Yellowhead County) 2023-06-09 18:22 to 06-15 18:05")):
    d = (pd.Timestamp(e) - pd.Timestamp(s)).days
    evr.append(dict(year=2023, place=place, prov="AB", people=POP(code), start=s, end=e, days_low=d, days=d, days_high=d, source=src + "; Census 2021"))
check(HRFE, "On the afternoon of May 28, 2023, within the boundaries of HRM, an outside fire ignited",
      "Evacuation ordered for Westwood Hills subdivision due to ongoing multiple structure and forest fire.", "All properties open to residents.")
check(os.path.join(R1, "ns", "governments-match-donations-red-cross-support-nova-scotians.html"), "many have returned home as some evacuation orders have been lifted")
NSU = os.path.join(R2, "ns3", "updates")
check(os.path.join(NSU, "ns_2023_05_29_update-wildfires-halifax-regional-municipality-shelburne-county.html"), "29 May 2023 | 12:49 PM",
      "Evacuation information is available at: https://www.facebook.com/shelburneemo")
check(os.path.join(NSU, "ns_2023_06_08_highway-103-shelburne-county-reopens-all-traffic.html"), "8 June 2023 | 5:17 PM",
      "A mandatory evacuation order remains in place for part of Shelburne County")
check(os.path.join(NSU, "ns_2023_06_09_highway-103-side-roads-reopen-access-properties.html"), "9 June 2023 | 12:27 PM",
      "Now that the mandatory evacuation orders have been lifted")
evr += [dict(year=2023, place="Upper Tantallon area (Halifax)", prov="NS", people=16400, start="2023-05-28", end="2023-06-06", days_low=9, days=9, days_high=17,
             source="HRFE Post Incident Analysis (HRM Regional Council 2024-08-06): ordered May 28, 'All properties open to residents' June 14 (high); NS release 2023-06-06: many returned as orders were lifted (central)"),
        dict(year=2023, place="Shelburne County (Barrington Lake fire)", prov="NS", people=7000, start="2023-05-29", end="2023-06-09", days_low=11, days=11, days_high=11,
             source="NS releases 2023-05-29 12:49 (evacuations under way), 06-08 17:17 (order remains for part of the county), 06-09 12:27 (orders lifted); HRFE report (nearly 7,000)")]
BCN = os.path.join(R2, "days2")
check(os.path.join(BCN, "bc_2023PREM0054-001323.html"), "Friday, August 18, 2023", "Thousands of people are under evacuation orders")
check(os.path.join(BCN, "bcnews", "n29393.html"), "As of 1 p.m. on Aug. 20, 2023, approximately 30,000 people are under an evacuation order and 36,000 are under an evacuation alert across B.C.")
check(os.path.join(BCN, "bcnews", "n29420.html"), "Saturday, August 26, 2023", "More than 15,000 people remain under an evacuation order")
check(os.path.join(BCN, "bc_2023EMCR0061-001387.html"), "Thursday, August 31, 2023", "In addition to 4,200 people on evacuation order")
check(os.path.join(BCN, "bc_2023EMCR0063-001463.html"), "As of Thursday, Sept. 14, about 370 people are on evacuation order")
BCPTS = [("2023-08-18", 30000), ("2023-08-20", 30000), ("2023-08-26", 15000), ("2023-08-31", 4200), ("2023-09-14", 370)]


def area(pts):
    return sum((pd.Timestamp(b[0]) - pd.Timestamp(a[0])).days * (a[1] + b[1]) / 2 for a, b in zip(pts, pts[1:]))


BC_PD = (area(BCPTS[1:]), area(BCPTS))        # low: no plateau before the first official count; central and high: 30,000 from Aug 18
evr.append(dict(year=2023, place="British Columbia, August surge (aggregate of all orders)", prov="BC", people=30000, start="2023-08-18", end="2023-09-14",
                days_low=BC_PD[0] / 30000, days=BC_PD[1] / 30000, days_high=BC_PD[1] / 30000,
                source="BC EMCR and Premier releases 2023-08-18, 08-20 (30,000 on order), 08-26 (15,000), 08-31 (4,200), 09-14 (370); days = person-days (straight lines between counts) / 30,000"))
check(os.path.join(R2, "days", "nl_0723n03.html"), "9,600 residents of Labrador City received notice to evacuate immediately")
check(os.path.join(R2, "days", "nl_0720n01.html"), "the evacuation order issued on Friday, July 12, 2024",
      "The evacuation order will be officially lifted for all residents of Labrador City, effective noon, Monday, July 22, 2024")
GZ = os.path.join(R1, "pages", "gazette_sor2025-10_ei_pilot23.html")
check(GZ, "approximately 2 400 people were evacuated from the fly-in community of Bunibonibee Cree Nation",
      "on July 22, 2024", "the evacuation order was lifted on August 17, 2024", "August 13, 2024")
check(os.path.join(R1, "nwt", "emar2025.txt"), "approximately 380 residents")
for place, prov, n, s, e, src in (("Jasper (residents)", "AB", POP("2021A00054815033"), "2024-07-22", "2024-08-17", "Canada Gazette SOR/2025-10 RIAS; Census 2021"),
                                  ("Labrador City", "NL", 9600, "2024-07-12", "2024-07-22", "NL releases 2024-07-20 and 07-23"),
                                  ("Bunibonibee Cree Nation", "MB", 2400, "2024-08-13", "2024-09-12", "Canada Gazette SOR/2025-10 RIAS"),
                                  ("Fort Good Hope", "NT", 380, "2024-06-15", "2024-07-06", "NWT Emergency Management Annual Report 2025")):
    d = (pd.Timestamp(e) - pd.Timestamp(s)).days
    evr.append(dict(year=2024, place=place, prov=prov, people=n, start=s, end=e, days_low=d, days=d, days_high=d, source=src))
evd = pd.DataFrame(evr)
evd.round({"days_low": 2, "days": 2, "days_high": 2}).to_csv(os.path.join(CF, "health2_days_away_events.csv"), index=False)
sub = evd.groupby("year").apply(lambda d: pd.Series({"people": d.people.sum(), "pd_low": (d.people * d.days_low).sum(),
                                                     "person_days": (d.people * d.days).sum(), "pd_high": (d.people * d.days_high).sum()}))
sub["mean_days"] = sub.person_days / sub.people
out23 = evd[(evd.year == 2023) & (evd.prov != "NT")]
lw = pd.read_csv(os.path.join(CF, "health_lost_wages_by_year.csv")).set_index("year")
check(os.path.join(R1, "ab", "scss_2023_24.txt"), "Over 37,000 wildfire evacuees were provided Emergency Evacuation Payments in 2023")
D_old = {2023: (7.0, 12.5, 18.0), 2024: (7.0, 12.5, 18.0)}
D_new = {}
for y in (2023, 2024):
    ppl = sub.loc[y, "people"]
    res = []
    for k, (col, pdc) in enumerate((("evacuees_low", "pd_low"), ("evacuees_central", "person_days"), ("evacuees_high", "pd_high"))):
        N = lw.loc[y, col]
        res.append((sub.loc[y, pdc] + (N - ppl) * D_old[y][k]) / N)
    D_new[y] = tuple(res)

FNF = float(lw.loc[2023, "fn_employment_factor"])
lwrows = []
for y in (2023, 2024, 2026):
    ER, AWE = float(lw.loc[y, "employed_per_person"]), float(lw.loc[y, "avg_weekly_earnings"])
    if y == 2026:
        N = n26; FN = (FH, 8805, 0); D = (7.0, 12.5, 18.0)
    else:
        N = tuple(float(lw.loc[y, c]) for c in ("evacuees_low", "evacuees_central", "evacuees_high"))
        FN = (float(lw.loc[y, "fn_on_reserve_adjusted_low"]), float(lw.loc[y, "fn_on_reserve_adjusted_central"]), 0.0)
        D = D_new[y]
    res = []
    for k in range(3):
        workers = (N[k] - FN[k]) * ER + FN[k] * ER * FNF
        wd = workers * D[k] * 5 / 7
        res.append((N[k], FN[k], workers, D[k], wd, wd * AWE / 5, to25(wd * AWE / 5, y)))
    lwrows.append(dict(year=y, **{f"{a}_{b}": round(res[k][i], 2) for i, a in enumerate(["evacuees", "fn_adjusted", "workers", "days_away", "workdays_lost", "wages_nominal", "wages_2025d"]) for k, b in enumerate(["low", "central", "high"])}))
    old_lw = H1[H1.item.eq("Lost work and wages of evacuated workers") & H1.period.eq(str(y))].iloc[0]
    if y in (2023, 2024):
        dn = (f"days away {D[0]:.1f} / {D[1]:.1f} / {D[2]:.1f} (was 7 / 12.5 / 18): official dates or counts for {int(sub.loc[y, 'people']):,} of {N[1]:,.0f} evacuees "
              f"({sub.loc[y, 'mean_days']:.1f} days on average, central)" + (f", including {int(out23.people.sum()):,} outside the NWT (Alberta, Nova Scotia, BC: {(out23.people * out23.days).sum() / out23.people.sum():.1f} days)" if y == 2023 else "")
              + "; the rest at the first-round 7 / 12.5 / 18")
    else:
        dn = f"evacuees {N[0]:,} / {N[1]:,} / {N[2]:,} (was 20,847: ISC 8,805 + Summerland 12,042); BC's count of more than 20,000 includes Summerland; days 7 / 12.5 / 18 unchanged"
    add(item="Lost work and wages of evacuated workers (revised)", jurisdiction="Canada", period=str(y),
        quantity=f"{res[1][4]:,.0f} workdays lost ({res[0][4]:,.0f} to {res[2][4]:,.0f})", unit_cost=f"SEPH average weekly earnings {AWE:.2f} / 5 per workday (nominal {y})",
        amount_low_cad=res[0][6], amount_central_cad=res[1][6], amount_high_cad=res[2][6],
        method="first-round formula unchanged (evacuees x employed per resident x days away x 5/7 x daily earnings; First Nations on-reserve evacuees at 0.636 of the employment ratio in low and central); only the inputs named in the notes change",
        source_name="first-round inputs (health_lost_wages_by_year.csv) + " + ("NWT TD 1007-19(2) chronology, Alberta Emergency Alerts, Nova Scotia releases and HRM post incident analysis, BC EMCR releases, NL releases, Canada Gazette SOR/2025-10, NWT EM Annual Report 2025, Census 2021" if y < 2026 else "BC EMCR 2026EMCR0037, ISC 2026, GNWT 2026 updates, Census 2021"),
        url="costs-full/health2_days_away_events.csv ; costs-full/health2_lost_wages_by_year.csv",
        overlaps_with=f"REPLACES health_costs.csv 'Lost work and wages of evacuated workers' {y} (central {float(old_lw.amount_central_cad) / 1e6:.1f} M): drop it when adding this; same overlaps with econ_ output rows and income transfers as that row",
        additive="yes", notes="ESTIMATE. " + dn + ".")
pd.DataFrame(lwrows).to_csv(os.path.join(CF, "health2_lost_wages_by_year.csv"), index=False)
check(os.path.join(R2, "qc", "qc_news_48601.html"), "plus de 15 000 personnes ont dû évacuer leur domicile")
nwt23 = evd[(evd.year == 2023) & (evd.prov == "NT")]
add(item="Average days away from home, evacuations with official dates (subset)", jurisdiction="Canada", period="2023 and 2024",
    quantity=(f"2023: {int(sub.loc[2023, 'people']):,} evacuees in {int((evd.year == 2023).sum())} events, {sub.loc[2023, 'mean_days']:.1f} days on average "
              f"(NWT {int(nwt23.people.sum()):,} in 12 events, {(nwt23.people * nwt23.days).sum() / nwt23.people.sum():.1f} days; Alberta: Drayton Valley 12, Fox Creek 18, Edson 3 and 6; "
              f"Nova Scotia: Upper Tantallon 9 (17 high), Shelburne County 11; BC August surge {BC_PD[1] / 30000:.1f} ({BC_PD[0] / 30000:.1f} low)); "
              f"2024: {int(sub.loc[2024, 'people']):,} evacuees (Jasper residents, Labrador City, Bunibonibee, Fort Good Hope), {sub.loc[2024, 'mean_days']:.1f} days; "
              "2025: Ontario's 8,387 wildfire evacuees hosted about one month; ESDC 2025 orders over 18 days on average"),
    dollars_basis="", method="evacuee-weighted mean of (end date - start date) for events with official start and end dates; BC: person-days from official counts of people on order at five dates; people = 2021 Census population or the official evacuee count",
    source_name="NWT TD 1007-19(2); Alberta Emergency Alerts archive (alberta.ca/aea/cap); Nova Scotia news releases 2023-05-29, 06-06, 06-08, 06-09; HRFE Post Incident Analysis (HRM, 2024); BC news releases 2023-08-18 to 09-14; NL releases 2024-07-20, 07-23; Canada Gazette SOR/2025-10; NWT EM Annual Report 2025; Ontario 2025 annual report; ESDC QP note EF_063_20260105",
    url="https://www.ntlegislativeassembly.ca/sites/default/files/legacy/td_1007-192.pdf ; https://www.alberta.ca/aea/cap/2023/05/16/2023-05-16T13_54_37-06_00=BrazeauCounty=013F2806-28ED-4853-B285-8D9FB419B304.htm ; https://www.halifax.ca/sites/default/files/documents/city-hall/regional-council/240806rci01.pdf ; https://news.gov.bc.ca/releases/2023EMCR0063-001463 ; https://www.gov.nl.ca/releases/2024/jps/0720n01/ ; https://gazette.gc.ca/rp-pr/p2/2025/2025-02-12/html/sor-dors10-eng.html",
    additive="no", notes="Quebec 2023 is still open (more than 15,000 evacuated by June 9, 2023 per the MSP release; 8,911 evacuation-benefit claims): no official end date was found in the MSP annual report 2023-24, the quebec.ca program releases of 2023-06-09 and 2024-06-23, or web searches for municipal and MSP notices on Chibougamau, Lebel-sur-Quevillon, Chapais and Sept-Iles (news reports only). Nova Scotia's dates and BC's August counts are official; the BC person-days use straight lines between the counts. The CWFEDB's event-level dates are not public (data-sharing agreement required).")

# EI: official count (2024) and excess initial claims from StatCan 14-10-0005 (2023 NWT, 2025 Manitoba)
ei = pd.read_csv(os.path.join(R2, "statcan", "t14100005", "14100005.csv"), usecols=["REF_DATE", "GEO", "Type of claim", "Claim detail", "VALUE"])
ei = ei[(ei["Type of claim"] == "Initial claims") & (ei["Claim detail"] == "Received")].copy()
ei["y"] = ei.REF_DATE.str[:4].astype(int); ei["m"] = ei.REF_DATE.str[5:7].astype(int)
piv = ei[ei.y.between(2019, 2026)].pivot_table(index=["GEO", "m"], columns="y", values="VALUE")
eir = []
nt = piv.loc["Northwest Territories"]
base_nt = nt[[2019, 2021, 2022, 2024, 2025]].loc[[8, 9]].mean(axis=1)
ex_nt = (nt.loc[[8, 9], 2023] - base_nt).sum()
eir.append(dict(case="NWT Aug-Sep 2023", actual=float(nt.loc[[8, 9], 2023].sum()), baseline=float(base_nt.sum()), excess=float(ex_nt)))
mb = piv.loc["Manitoba"]
base_mb = mb.loc[6, [2019, 2021, 2022, 2023, 2024]].mean()
prov_other = ["Alberta", "British Columbia", "Quebec", "Ontario", "Saskatchewan", "Nova Scotia", "New Brunswick", "Newfoundland and Labrador"]
rise = np.mean([piv.loc[g].loc[6, 2025] / piv.loc[g].loc[6, [2019, 2021, 2022, 2023, 2024]].mean() for g in prov_other])
ex_mb_hi = mb.loc[6, 2025] - base_mb
ex_mb_lo = mb.loc[6, 2025] - base_mb * rise
eir.append(dict(case="Manitoba June 2025", actual=float(mb.loc[6, 2025]), baseline=float(base_mb), excess=float(ex_mb_hi),
                excess_after_national_rise=float(ex_mb_lo), other_provinces_june_rise=float(rise)))
pd.DataFrame(eir).round(1).to_csv(os.path.join(CF, "health2_ei_excess_claims.csv"), index=False)
add(item="EI claims tied to wildfire evacuations (counts, no dollars)", jurisdiction="Canada", period="2023 to 2025",
    quantity=f"2024: over 3,054 EI applications related to the 2024 wildfire season by 2024-09-09 (ESDC), about half from Jasper; 2023 NWT: about {ex_nt:,.0f} extra initial claims in Aug-Sep 2023; 2025 Manitoba: {ex_mb_lo:,.0f} to {ex_mb_hi:,.0f} extra initial claims in June 2025",
    dollars_basis="", method="official ESDC counts where published; otherwise ESTIMATE = initial claims received (StatCan 14-10-0005, unadjusted) minus the same months' average of other years (2019, 2021, 2022, 2024, 2025 for the NWT; 2019, 2021-2024 for Manitoba), Manitoba's lower bound also removes the June 2025 rise seen in eight other provinces",
    source_name="ESDC QP note CSDec2024_009; Canada Gazette SOR/2025-10 RIAS; Statistics Canada Table 14-10-0005-01 (EI claims received, by province and territory, monthly)",
    url="https://search.open.canada.ca/qpnotes/record/esdc-edsc,CSDec2024_009 ; https://www150.statcan.gc.ca/n1/tbl/csv/14100005-eng.zip",
    additive="no", notes=f"Counts only (EI benefits are transfers that replace part of the lost wages). No clear wildfire signal in provincial EI claims for AB 2023, NS 2023, BC 2023 or NL 2024. The June 2025 rise in eight other provinces averaged {100 * (rise - 1):.0f}% (national labour market, not fires).")
add(item="Official count of work hours or days lost to wildfire evacuations", jurisdiction="Canada", period="2023 to 2026", quantity="none published",
    dollars_basis="", method="search", source_name="Statistics Canada LFS releases (The Daily, May 2023 to Aug 2026, 12 files); StatCan subject page 'Workplace health and work absences'; LFS tables of absence by reason (none by province-month and reason 'other'); NWT Bureau of Statistics monthly LFS pages (404)",
    url="https://www150.statcan.gc.ca/n1/en/subjects/labour/workplace_health_and_work_absences", additive="no",
    notes="LFS mentions wildfires in text only; reason-for-absence detail by province and month is not in public tables.")

# =================================================================== GAP 5: HEALTH SYSTEM COSTS OF EVACUATIONS
check(os.path.join(AS471, "1104-2686.txt"), "l'établissement a procédé à une réclamation d'un montant de 806 421$")
check(os.path.join(AS471, "1104-5127.txt"), "MSSS-Feux de forêt 338 713 $")
check(os.path.join(AS471, "1104-5234.txt"), "Feux de forêt 34 712 687 712 687")
qh = 806421 + 338713 + 712687
add(item="Health care (health system costs of evacuations): Quebec health establishments' extra costs of the 2023 forest-fire emergency", jurisdiction="QC", period="2023 (FY2023-24)",
    quantity="CRSSS de la Baie-James 806,421 (claim for the emergency, towns evacuated); CIUSSS du Saguenay-Lac-Saint-Jean 338,713 (MSSS 'feux de forêt' funding); CISSS de la Côte-Nord 712,687 ('Feux de forêt')",
    amount_low_cad=to25(qh, 2023), amount_central_cad=to25(qh, 2023), amount_high_cad=to25(qh, 2023),
    method="amounts in the audited 2023-24 AS-471 financial reports of the three establishments", source_name="MSSS, Rapports financiers annuels AS-471 2023-2024 (codes 1104-2686, 1104-5127, 1104-5234)",
    url="https://publications.msss.gouv.qc.ca/msss/fichiers/statistiques/rapports-financiers/2024/1104-2686.pdf ; .../1104-5127.pdf ; .../1104-5234.pdf",
    overlaps_with="DFAA (the Baie-James claim goes to the federal government through the MSSS); not in health_ or gov_", additive="yes",
    notes="Only the 25 AS-471 reports downloaded were searched; MUHC, CHU Sainte-Justine, IUCPQ, Nunavik and Cree Board reports were not found.")
ab = pd.read_csv(os.path.join(R2, "ab", "grants_2023-24.csv"), low_memory=False, encoding="latin-1")
ab.columns = [c.strip() for c in ab.columns]
ahs = ab[(ab.Recipient == "ALBERTA HEALTH SERVICES") & ab.Program.isin(["2023 NWT EVACUATION GOARI", "2023 NWAB WILDFIRE"])]
nwt_ahs = float(ahs[ahs.Program == "2023 NWT EVACUATION GOARI"].Amount.sum())
nwab_ahs = float(ahs[ahs.Program == "2023 NWAB WILDFIRE"].Amount.sum())
assert abs(nwt_ahs - 521552.86) < 0.01 and abs(nwab_ahs - 157847.22) < 0.01
add(item="Health care (health system costs of evacuations): Alberta Health Services costs reimbursed for the 2023 NWT evacuation", jurisdiction="AB (for NT)",
    period="2023 (FY2023-24)", quantity="2 payments, 2024-03-21", amount_low_cad=to25(nwt_ahs, 2023), amount_central_cad=to25(nwt_ahs, 2023), amount_high_cad=to25(nwt_ahs, 2023),
    method="official grant payments (Public Safety and Emergency Services, program '2023 NWT EVACUATION GOARI')", source_name="Government of Alberta, Grant payments disclosure 2023-24",
    url="https://open.alberta.ca/dataset/faef5592-9ea0-4b60-92db-79a8a9673fe5/resource/0c7e93c1-1c28-460d-92eb-f6934e9ba4a8/download/tbf-grants-disclosure-2023-24.csv",
    overlaps_with="inside Alberta's $44.3M NWT hosting row (health_costs.csv, additive)", additive="no", notes="Reference: counted in the hosting row.")
add(item="Health care (health system costs of evacuations): Alberta Health Services costs reimbursed for the 2023 northwest Alberta wildfire", jurisdiction="AB", period="2023 (FY2023-24)",
    quantity="1 payment, 2024-01-29", amount_low_cad=to25(nwab_ahs, 2023), amount_central_cad=to25(nwab_ahs, 2023), amount_high_cad=to25(nwab_ahs, 2023),
    method="official grant payment (Disaster Recovery Program '2023 NWAB WILDFIRE')", source_name="Government of Alberta, Grant payments disclosure 2023-24",
    url="https://open.alberta.ca/dataset/faef5592-9ea0-4b60-92db-79a8a9673fe5/resource/0c7e93c1-1c28-460d-92eb-f6934e9ba4a8/download/tbf-grants-disclosure-2023-24.csv",
    overlaps_with="inside the first-round high-case Alberta 2023 DRP remainder (high case only)", additive="yes",
    notes="Counted in low and central here; the high case already holds it (small overlap, 0.16 M).")
add(item="Health care (health system and mental health costs of evacuations or smoke): other jurisdictions", jurisdiction="Canada", period="2023 to 2026",
    quantity="not quantified in any official record found", dollars_basis="", method="search",
    source_name="NWT Supplementary Estimates (Operations) No. 2 and No. 4 2023-24 (no Health and Social Services wildfire item); NWT 2023 Wildfire Response Review; Manitoba 2025 Wildfire Interim After Action Review (mental health supports described, no dollars); ISC QP note ISC-2026-QP-00844 (mental wellness supports inside EMAP, no dollars); Public Safety Canada 2026 updates; Alberta grant payments 2023-24 to 2025-26 ('2023 AB SPRING WILDFIRES - GOARI' 0.92 M and 1.26 M with recipients withheld)",
    url="https://www.ntlegislativeassembly.ca/sites/default/files/legacy/td_29-201.pdf", additive="no",
    notes="ISC funds evacuee mental wellness supports inside EMAP (in gov_). No provincial health authority report with wildfire mental-health spending was found.")

# =================================================================== write
df = pd.DataFrame(rows)[COLS]
txt = " ".join(df.astype(str).values.ravel())
assert "\u2014" not in txt, "em dash in output"
df.to_csv(os.path.join(CF, "health2_costs.csv"), index=False)

# summary: second-round additive totals and net bill after replacements
a = df[df.additive.isin(["yes", "partial"])].copy()
for c in ("amount_low_cad", "amount_central_cad", "amount_high_cad"):
    a[c] = pd.to_numeric(a[c])
a["group"] = np.where(a.item.str.startswith("Health care"), "Health care", np.where(a.item.str.startswith("Lost work"), "Lost work and wages", "Evacuations"))
s = a.groupby("group")[["amount_low_cad", "amount_central_cad", "amount_high_cad"]].sum()
s.loc["TOTAL"] = s.sum()
print("health2 additive (incl. replacements, before removing replaced first-round rows), $M:")
print((s / 1e6).round(2).to_string())
rep = H1[H1.item.str.startswith("Health care: smoke-related emergency visits and admissions, scaled")
         | (H1.item.eq("Lost work and wages of evacuated workers") & H1.period.isin(["2023", "2024", "2026"]))
         | H1.item.str.startswith("Evacuations: NWT Evacuation Travel Support Program")]
r = rep[["amount_low_cad", "amount_central_cad", "amount_high_cad"]].astype(float).sum()
print("first-round rows replaced, $M:", (r / 1e6).round(2).to_dict())
h1a = H1[H1.additive == "yes"][["amount_low_cad", "amount_central_cad", "amount_high_cad"]].astype(float).sum()
print("first-round additive people costs $M:", (h1a / 1e6).round(1).to_dict())
print("new people-cost total $M:", ((h1a - r + s.loc["TOTAL"].values) / 1e6).round(1).to_dict())
print(tot[["hosp_ed_central", "hosp_adm_central", "phys_central"]].round(0))
print("days", {y: tuple(round(x, 2) for x in D_new[y]) for y in D_new}, sub.round(2).to_dict())
print("QC ED", round(QC_ED, 2), "visits", QC_VIS, "ISC per evacuee", {k: round(v) for k, v in ISC_PER.items()}, "per_bc", round(per_bc), "EI", eir)
print("rows", len(df))
