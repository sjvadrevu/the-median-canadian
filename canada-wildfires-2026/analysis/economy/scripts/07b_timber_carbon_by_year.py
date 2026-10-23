"""Per-year series 2016-2026 (plus 2023-2026 union) of burned area, forest, timber volume, biomass and carbon, set
against that year's harvest and national emissions.

Pre-fire SCANFI v2 biomass layer by fire year: 2016-2020 -> 2015 map, 2021-2025 -> 2020 map, 2026 -> 2025 map.
Volume and the combustion shares: NFI kNN 2011 (as in 07). Same carbon assumptions as 07 (0.5 C fraction; low =
foliage share, high = foliage + branch + bark share; CO2 = C x 44/12; aboveground live biomass only).
Comparators per year (latest available year used when the fire year is newer, flagged in *_year columns):
  NFD net merchantable roundwood harvest (to 2024); NIR 2026 national total and passenger cars/light trucks/
  motorcycles (to 2024); NIR Table 6-5 official "wildfires - direct immediate emissions" in the MANAGED forest
  (all carbon pools, CO2+CH4+N2O, to 2024) shown for context, not as a like-for-like check.
Scaled = per-cell provincial factor official/mapped for that year (work/scale_factors_2016_2025.csv; 2026 factors from
analysis/forest; national parks unscaled before 2026, Parks Canada factor in 2026).
Outputs: timber_carbon_by_year.csv (Canada and province rows), work/log_07b.txt
"""
import numpy as np
import pandas as pd
import openpyxl
from ecommon import WORK, OUT, RAW

c26 = pd.read_pickle(WORK / "cells.pkl")
c16 = pd.read_pickle(WORK / "cells_1622.pkl")
c26["bio2020"] = np.load(WORK / "scanfi_bio2020.npz")["bio_t"]
c26["bio2025"] = np.load(WORK / "scanfi_bio2025.npz")["bio_t"]
c16["bio2015"] = np.load(WORK / "scanfi_bio2015_1622.npz")["bio_t"]
c16["bio2020"] = np.load(WORK / "scanfi_bio2020_1622.npz")["bio_t"]
K = ["vol_merch", "vol_total", "agb_live", "bio_foliage", "bio_branch", "bio_stembark"]
for df, f in [(c26, "knn2011.npz"), (c16, "knn2011_cells_1622.npz")]:
    k = np.load(WORK / f)
    for v in K:
        df[v] = k[v + "_per_ha"].astype("float64") * df.area_ha.values
sf = pd.read_csv(WORK / "scale_factors_2016_2025.csv")
fac = {(r.year, r.prov): r.factor for r in sf.itertuples()}

h = pd.read_csv(RAW / "nfd" / "NFD - Net Merchantable Volume of Roundwood Harvested by Category and Ownership - EN FR.csv",
                encoding="utf-8-sig")
hv = h.groupby(["ISO", "Year"])["Volume (cubic metres)"].sum().unstack()
hv.loc["Canada"] = hv.sum()
g = pd.read_csv(RAW / "ghg" / "B-Economic-Sector_GHG_Econ_Can_Prov_Terr.csv", encoding="latin-1")
g = g[g.Region == "Canada"].copy()
g["Total (kt CO2e)"] = pd.to_numeric(g["Total (kt CO2e)"], errors="coerce")
nat = g[g.Index == 0].set_index("Year")["Total (kt CO2e)"] * 1000
cars = g[g.Index == 19].set_index("Year")["Total (kt CO2e)"] * 1000
wb = openpyxl.load_workbook(RAW / "ghg" / "E-LULUCF_EN_Ch6_Tables_FullTimeSeries.xlsx", data_only=True)
ws = wb["Table 6–5 All Years"]
yrs = [c.value for c in ws[3]][4:]
wf = {int(y): ws.cell(row=17, column=5 + i).value for i, y in enumerate(yrs) if y is not None}
wfco2 = {int(y): ws.cell(row=18, column=5 + i).value for i, y in enumerate(yrs) if y is not None}

rows = []
for y in list(range(2016, 2027)) + ["2023_2026"]:
    if y == "2023_2026":
        d = c26.copy()
        first = np.where(d.b2023, 2023, np.where(d.b2024, 2024, np.where(d.b2025, 2025, 2026)))
        d["bio"] = np.where(first == 2026, d.bio2025, d.bio2020)
        d["w"] = d.w_last
    else:
        src = c16 if y < 2023 else c26
        d = src[src[f"b{y}"]].copy()
        d["bio"] = d.bio2015 if y <= 2020 else (d.bio2020 if y <= 2025 else d.bio2025)
        if y == 2026:
            d["w"] = d.w26
        else:
            d["w"] = [1.0 if pk > 0 else fac.get((y, p), np.nan) for p, pk in zip(d.prov, d.park_id)]
            d["w"] = d.w.fillna(1.0)
    for p, gg in list(d.groupby("prov")) + [("Canada", d)]:
        fol = gg.bio_foliage.sum() / gg.agb_live.sum()
        nst = (gg.bio_foliage.sum() + gg.bio_branch.sum() + gg.bio_stembark.sum()) / gg.agb_live.sum()
        r = dict(year=y, prov=p, area_mapped_ha=gg.area_ha.sum(), area_scaled_ha=(gg.area_ha * gg.w).sum(),
                 forest_mapped_ha=gg.forest_ha.sum(), agb_prefire_t=gg.bio.sum(), agb_prefire_scaled_t=(gg.bio * gg.w).sum(),
                 vol_merch_m3=gg.vol_merch.sum(), vol_merch_scaled_m3=(gg.vol_merch * gg.w).sum(),
                 vol_total_m3=gg.vol_total.sum(), foliage_share=fol, nonstem_share=nst)
        for s in ["", "_scaled"]:
            b = r[f"agb_prefire{s}_t"]
            r[f"co2_low{s}_t"] = b * 0.5 * fol * 44 / 12
            r[f"co2_high{s}_t"] = b * 0.5 * nst * 44 / 12
        hy = y if (y != "2023_2026" and y <= 2024) else 2024
        if p in hv.index:
            r["harvest_year"] = hy
            r["harvest_m3"] = hv.loc[p, hy] if y != "2023_2026" else hv.loc[p, 2024]
            r["years_of_harvest_scaled"] = r["vol_merch_scaled_m3"] / r["harvest_m3"]
            r["years_of_harvest_mapped"] = r["vol_merch_m3"] / r["harvest_m3"]
        if p == "Canada":
            ey = y if (y != "2023_2026" and y <= 2024) else 2024
            r["emissions_year"] = ey
            r["national_total_t"] = nat[ey]
            r["cars_lighttrucks_moto_t"] = cars[ey]
            for s in ["", "_scaled"]:
                r[f"co2_low{s}_share_national"] = r[f"co2_low{s}_t"] / nat[ey]
                r[f"co2_high{s}_share_national"] = r[f"co2_high{s}_t"] / nat[ey]
            if y != "2023_2026" and y <= 2024:
                r["nir_managed_forest_wildfire_direct_immediate_t_co2e"] = wf[y] * 1000
                r["nir_managed_forest_wildfire_direct_immediate_t_co2"] = wfco2[y] * 1000
        rows.append(r)
o = pd.DataFrame(rows)
o.to_csv(OUT / "timber_carbon_by_year.csv", index=False)
pd.set_option("display.width", 250)
cn = o[o.prov == "Canada"]
print(cn[["year", "area_mapped_ha", "area_scaled_ha", "agb_prefire_t", "vol_merch_scaled_m3", "harvest_m3",
          "years_of_harvest_scaled", "co2_low_scaled_t", "co2_high_scaled_t", "national_total_t",
          "co2_low_scaled_share_national", "co2_high_scaled_share_national",
          "nir_managed_forest_wildfire_direct_immediate_t_co2e"]].to_string())
